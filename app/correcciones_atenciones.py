"""Correcciones desde el Excel original de Cargue Atenciones, con revision y auditoria."""
from datetime import datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import json
from uuid import uuid4

from flask import current_app, jsonify, request, send_file
from flask_login import login_required
from itsdangerous import URLSafeTimedSerializer, BadSignature
from openpyxl import load_workbook, Workbook
from sqlalchemy import or_
from werkzeug.exceptions import RequestEntityTooLarge

from app.models import db, AtencionDiaDetalle, AuditLog, PrefacturaComercial, Usuario
from app.routes import comercial_bp

CAMPOS = {
    'nro_orden': 'Orden', 'nro_factura': 'Factura', 'fecha_factura': 'Fecha factura',
    'precio': 'Precio', 'forma_pago': 'Forma de pago', 'servicio': 'Examen',
    'nro_identificacion': 'Documento paciente', 'nombre_paciente': 'Paciente',
    'acuerdo_comercial': 'Acuerdo comercial', 'empresa_mision': 'Empresa en mision',
    'sede': 'Sede', 'nombre_vendedor': 'Vendedor', 'fecha_creacion_orden': 'Fecha atencion',
    'usuario_creacion': 'Usuario creacion', 'estado_orden': 'Estado orden',
    'fecha_anulacion': 'Fecha anulacion', 'cliente_id': 'Empresa vinculada', 'vendedor_id': 'Vendedor vinculado',
}
CAMPOS_EXCEL = list(CAMPOS)[:16]


def snapshot(reg):
    def valor(v):
        if isinstance(v, datetime):
            return v.isoformat()
        if isinstance(v, Decimal):
            return str(v)
        return v
    return {c.name: valor(getattr(reg, c.name)) for c in reg.__table__.columns}


def revision():
    return URLSafeTimedSerializer(current_app.config['SECRET_KEY'], salt='revision-correcciones-v2')


def _comprobar_editable(reg, scope, fecha=None):
    from app.routes import cargue_atenciones as ca
    ca._asegurar_acceso_registro_atencion(reg, scope)
    if (reg.estado_gestion or '').upper() != 'CARGADA':
        raise ValueError('La atencion ya no esta en estado CARGADA')
    empresa = (PrefacturaComercial.cliente_id == reg.cliente_id if reg.cliente_id is not None
               else PrefacturaComercial.nombre_empresa == (reg.acuerdo_comercial or reg.empresa_mision or 'SIN_EMPRESA'))
    bloqueada = PrefacturaComercial.query.filter(
        empresa, PrefacturaComercial.fecha_desde <= (fecha or reg.fecha_creacion_orden),
        PrefacturaComercial.fecha_hasta >= (fecha or reg.fecha_creacion_orden),
        or_(PrefacturaComercial.estado == 'CERRADA', PrefacturaComercial.bloqueada_por_pago.is_(True),
            PrefacturaComercial.pagos_cartera.any()),
    ).first()
    if bloqueada:
        raise ValueError('La atencion pertenece a una prefactura cerrada o con pagos')


def _leer_archivos(archivos, scope, bloquear=False):
    from app.routes import cargue_atenciones as ca
    from types import SimpleNamespace
    try:
        desde = datetime.strptime(request.form.get('periodo_desde', ''), '%Y-%m-%d')
        hasta = datetime.strptime(request.form.get('periodo_hasta', ''), '%Y-%m-%d').replace(hour=23, minute=59, second=59, microsecond=999999)
        if desde > hasta:
            raise ValueError
    except ValueError:
        raise ValueError('Selecciona el periodo de las atenciones que vas a corregir')
    cambios, vistos, huellas = [], set(), []
    clientes = ca._construir_lookup_clientes()
    vendedores = ca._construir_lookup_vendedores()
    for archivo in archivos:
        nombre = (archivo.filename or '').replace('\\', '/').split('/')[-1]
        contenido = archivo.read(20 * 1024 * 1024 + 1)
        if len(contenido) > 20 * 1024 * 1024:
            raise ValueError(f'{nombre}: el archivo supera 20 MB')
        huellas.append([nombre, sha256(contenido).hexdigest()])
        filas = ca._leer_excel_atenciones(contenido, nombre)
        # Validar el mismo encabezado y aliases que usa Cargue Atenciones.
        ca._extraer_registros_excel(filas[:1])
        if len(filas) > 50001:
            raise ValueError(f'{nombre}: maximo 50000 filas por archivo')
        for numero, fila in enumerate(filas[1:], 2):
            if not fila or all(v in (None, '') for v in fila):
                continue
            try:
                registro = ca._extraer_registros_excel([filas[0], fila])[0]
                valores = {}
                for campo, encabezado in zip(CAMPOS_EXCEL, ca.COLUMNAS_ESPERADAS):
                    valor = registro.get(encabezado)
                    if campo.startswith('fecha_'):
                        parsed = ca._parse_fecha(valor)
                        if valor not in (None, '') and parsed is None:
                            raise ValueError(f'{CAMPOS[campo]} no es una fecha valida')
                        valores[campo] = parsed
                    elif campo == 'precio':
                        parsed = ca._parse_precio(valor)
                        if parsed is None or not parsed.is_finite() or parsed < 0 or parsed >= Decimal('10000000000000') or parsed != parsed.quantize(Decimal('.01')):
                            raise ValueError('Precio invalido; usa un valor no negativo con hasta dos decimales')
                        valores[campo] = parsed.quantize(Decimal('.01'))
                    else:
                        if isinstance(valor, float) and valor.is_integer():
                            valor = int(valor)
                        valores[campo] = ca._normalizar(valor)
                        limite = AtencionDiaDetalle.__table__.columns[campo].type.length
                        if valores[campo] and len(valores[campo]) > limite:
                            raise ValueError(f'{CAMPOS[campo]} admite hasta {limite} caracteres')
                for campo in ('nro_orden', 'servicio', 'nro_identificacion', 'nombre_paciente', 'fecha_creacion_orden'):
                    if not valores[campo]:
                        raise ValueError(f'{CAMPOS[campo]} es obligatorio')
                if not desde <= valores['fecha_creacion_orden'] <= hasta:
                    raise ValueError('La fecha de atencion esta fuera del periodo seleccionado')
                query = AtencionDiaDetalle.query.filter(
                    AtencionDiaDetalle.nro_orden == valores['nro_orden'],
                    AtencionDiaDetalle.fecha_creacion_orden >= desde,
                    AtencionDiaDetalle.fecha_creacion_orden <= hasta)
                if not ca._is_admin_user():
                    query = query.filter(ca._condicion_scope_atenciones(scope))
                candidatos = (query.with_for_update() if bloquear else query).all()
                exactos = [r for r in candidatos if r.servicio == valores['servicio'] and r.nro_identificacion == valores['nro_identificacion']]
                # Si se corrige documento o examen, el otro dato debe identificar una sola fila.
                posibles = exactos or [r for r in candidatos if r.servicio == valores['servicio'] or r.nro_identificacion == valores['nro_identificacion']]
                if len(posibles) != 1:
                    raise ValueError('No se encontro una atencion unica por orden, examen y documento. Revisa la coincidencia en Consulta Atenciones; no se crean registros nuevos')
                reg = posibles[0]
                ca._asegurar_acceso_registro_atencion(reg, scope)
                if reg.id in vistos:
                    raise ValueError('La misma atencion aparece mas de una vez en los archivos')
                vistos.add(reg.id)
                original = snapshot(reg)
                if (valores['estado_orden'] or '').upper() == 'ANULADA' and valores['fecha_anulacion'] is None:
                    valores['fecha_anulacion'] = reg.fecha_anulacion or valores['fecha_creacion_orden']
                empresa_anterior = reg.cliente.razon_social if reg.cliente else reg.acuerdo_comercial or reg.empresa_mision or 'SIN_EMPRESA'
                # Reasociar empresa/vendedor solo cuando se corrigen sus columnas.
                if valores['acuerdo_comercial'] != reg.acuerdo_comercial or valores['empresa_mision'] != reg.empresa_mision:
                    cliente = ca._cliente_desde_registro(registro, clientes)
                    if cliente is None:
                        raise ValueError('El acuerdo comercial no identifica una empresa configurada de forma unica')
                    ca.exigir_cliente(cliente.id)
                    valores['cliente_id'] = cliente.id
                    valores['vendedor_id'] = cliente.vendedor_id
                if valores['nombre_vendedor'] != reg.nombre_vendedor and 'cliente_id' not in valores:
                    vendedor = reg.cliente.vendedor if reg.cliente else ca._vendedor_desde_registro(registro, vendedores)
                    if vendedor is None and valores['nombre_vendedor']:
                        raise ValueError('El vendedor no se pudo identificar')
                    valores['vendedor_id'] = vendedor.id if vendedor else None
                cambios_fila = {campo: valor for campo, valor in valores.items() if valor != getattr(reg, campo)}
                if not cambios_fila:
                    continue
                _comprobar_editable(reg, scope)
                destino = SimpleNamespace(**{c.name: getattr(reg, c.name) for c in reg.__table__.columns})
                for campo, valor in cambios_fila.items():
                    setattr(destino, campo, valor)
                _comprobar_editable(destino, scope, destino.fecha_creacion_orden)
                cambios.append((reg, cambios_fila, original, nombre, {
                    'desde': desde.strftime('%Y-%m-%d'), 'hasta': hasta.strftime('%Y-%m-%d'),
                    'empresa_anterior': empresa_anterior}))
            except (ValueError, InvalidOperation) as exc:
                raise ValueError(f'{nombre}, fila {numero}: {exc}')
    return cambios, huellas


def _resumen(cambios):
    salida = []
    for reg, valores, original, archivo, firmado in cambios:
        for campo, valor in valores.items():
            salida.append({'atencion_id': reg.id, 'empresa': reg.cliente.razon_social if reg.cliente else reg.acuerdo_comercial,
                           'archivo': archivo, 'campo': CAMPOS[campo], 'antes': original[campo],
                           'despues': valor.isoformat() if isinstance(valor, datetime) else str(valor) if valor is not None else None})
    return salida


def _actualizar_borradores(cambios):
    """Mantener los totales persistidos consistentes, incluso al anular la ultima fila."""
    from app.routes import cargue_atenciones as ca
    catalogo = ca._construir_lookup_catalogo()
    revisadas = set()
    for reg, _, original, _, contexto in cambios:
        empresa = reg.cliente.razon_social if reg.cliente else reg.acuerdo_comercial or reg.empresa_mision or 'SIN_EMPRESA'
        fechas = [datetime.fromisoformat(original['fecha_creacion_orden']), reg.fecha_creacion_orden]
        prefs = PrefacturaComercial.query.filter(
            PrefacturaComercial.nombre_empresa.in_([empresa, contexto['empresa_anterior']]), PrefacturaComercial.estado == 'BORRADOR',
            PrefacturaComercial.origen == 'ATENCIONES',
            PrefacturaComercial.fecha_desde <= max(fechas), PrefacturaComercial.fecha_hasta >= min(fechas),
        ).with_for_update().all()
        for pref in prefs:
            if pref.id in revisadas:
                continue
            revisadas.add(pref.id)
            empresa_pref = pref.nombre_empresa
            duplicadas = PrefacturaComercial.query.filter_by(
                nombre_empresa=empresa_pref, fecha_desde=pref.fecha_desde, fecha_hasta=pref.fecha_hasta,
                origen='ATENCIONES').count()
            if duplicadas > 1:
                raise ValueError(f'{empresa}: hay varias prefacturas del mismo periodo; revisalas antes de corregir')
            query = AtencionDiaDetalle.query.filter(
                AtencionDiaDetalle.fecha_creacion_orden >= pref.fecha_desde,
                AtencionDiaDetalle.fecha_creacion_orden <= pref.fecha_hasta)
            if pref.cliente_id is not None:
                query = query.filter(AtencionDiaDetalle.cliente_id == pref.cliente_id)
            else:
                query = query.filter(AtencionDiaDetalle.cliente_id.is_(None))
            elegibles, formas = [], set()
            for fila in query.all():
                nombre = fila.cliente.razon_social if fila.cliente else fila.acuerdo_comercial or fila.empresa_mision or 'SIN_EMPRESA'
                pago = ca._normalizar_forma_pago(fila.forma_pago)
                if nombre != empresa_pref or (fila.estado_orden or '').strip().upper() == 'ANULADA':
                    continue
                if pago not in ('CREDITO', 'CONTADO', 'EFECTIVO', 'PARTICULAR', 'PARTICULARES'):
                    continue
                if ca._clasificar_servicio(fila.servicio, catalogo) == 'ECOBABY' or 'ECOBABY' in (fila.servicio or '').upper():
                    continue
                elegibles.append(fila)
                formas.add('CREDITO' if pago == 'CREDITO' else 'EFECTIVO')
            pref.valor_total = sum((fila.precio or Decimal(0) for fila in elegibles), Decimal(0))
            pref.cant_pacientes = len({(fila.nro_identificacion or '', fila.nombre_paciente or '') for fila in elegibles})
            if formas:
                pref.forma_pago = 'MIXTO' if len(formas) > 1 else next(iter(formas))
            pref.usuario_genera_id = ca.current_user.id


@comercial_bp.route('/atenciones-dia/correcciones-excel', methods=['POST'])
@login_required
def cargar_correcciones_excel():
    from app.routes import cargue_atenciones as ca
    try:
        ca._require_commercial_permission(ca.PERMISO_EDICION_ATENCIONES)
        scope = ca._resolver_vendedor_usuario_actual()
        if not ca._is_admin_user() and scope is None:
            raise PermissionError('No tienes un vendedor asociado')
        archivos = request.files.getlist('archivos')
        if not archivos or len(archivos) > 20:
            raise ValueError('Selecciona entre 1 y 20 archivos Excel')
        aplicar = request.form.get('accion', 'revisar') == 'aplicar'
        cambios, huellas = _leer_archivos(archivos, scope, bloquear=aplicar)
        resumen = _resumen(cambios)
        sello = {'usuario': ca.current_user.id, 'archivos': huellas,
                 'cambios': sha256(json.dumps({'resumen': resumen, 'originales': [c[2] for c in cambios], 'periodo': [request.form.get('periodo_desde'), request.form.get('periodo_hasta')]}, sort_keys=True).encode()).hexdigest()}
        if not aplicar:
            return jsonify(cambios=resumen, atenciones=len(cambios), token=revision().dumps(sello))
        try:
            aprobado = revision().loads(request.form.get('token', ''), max_age=3600)
        except BadSignature:
            raise ValueError('Revisa los archivos de nuevo antes de aplicar los cambios')
        if aprobado != sello:
            raise ValueError('Los archivos o las atenciones cambiaron; revisa nuevamente')
        lote, periodos = str(uuid4()), {}
        for reg, valores, original, archivo, firmado in cambios:
            ca._desvincular_prefacturas_de_atencion(reg)
            for campo, valor in valores.items():
                setattr(reg, campo, valor)
            db.session.flush()
            if 'cliente_id' in valores:
                db.session.expire(reg, ['cliente'])
            ca._sincronizar_cruce_prefacturas_con_atencion(reg)
            nuevo = snapshot(reg)
            empresa = reg.cliente.razon_social if reg.cliente else reg.acuerdo_comercial or reg.empresa_mision or 'SIN_EMPRESA'
            nuevo['_correccion'] = {'archivo': archivo, 'lote': lote, 'empresa': empresa,
                                     'usuario': str(getattr(ca.current_user, 'usuario', ''))}
            db.session.add(AuditLog(usuario_id=ca.current_user.id, tabla=AtencionDiaDetalle.__tablename__,
                                   registro_id=reg.id, accion='CORRECCION_EXCEL', datos_anteriores=original,
                                   datos_nuevos=nuevo, ip_address=request.remote_addr))
            if firmado['empresa_anterior'] != empresa:
                periodos[(firmado['empresa_anterior'], firmado['desde'], firmado['hasta'])] = {'empresa': firmado['empresa_anterior'], 'cliente_id': original['cliente_id'], 'fecha_desde': firmado['desde'], 'fecha_hasta': firmado['hasta']}
            key = (empresa, firmado['desde'], firmado['hasta'])
            periodos[key] = {'empresa': empresa, 'cliente_id': reg.cliente_id,
                            'fecha_desde': firmado['desde'], 'fecha_hasta': firmado['hasta']}
        _actualizar_borradores(cambios)
        db.session.commit()
        return jsonify(actualizadas=len(cambios), lote=lote, periodos=list(periodos.values()))
    except RequestEntityTooLarge:
        db.session.rollback()
        return jsonify(error='Los archivos superan el limite de carga. Divide la carga en grupos mas pequenos'), 413
    except (ValueError, PermissionError) as exc:
        db.session.rollback()
        return jsonify(error=str(exc)), 403 if isinstance(exc, PermissionError) else 400
    except Exception:
        db.session.rollback()
        current_app.logger.exception('Error aplicando correcciones Excel')
        return jsonify(error='No se pudieron guardar las correcciones; no se aplicaron cambios'), 500


@comercial_bp.route('/atenciones-dia/correcciones-informe', methods=['GET'])
@login_required
def informe_correcciones():
    from app.routes import cargue_atenciones as ca
    try:
        ca._require_commercial_permission(ca.PERMISO_CONSULTA_ATENCIONES)
        scope = ca._resolver_vendedor_usuario_actual()
        if not ca._is_admin_user() and scope is None:
            raise PermissionError('No tienes un vendedor asociado')
        query = db.session.query(AuditLog, AtencionDiaDetalle, Usuario).outerjoin(
            AtencionDiaDetalle, AtencionDiaDetalle.id == AuditLog.registro_id
        ).outerjoin(Usuario, Usuario.id == AuditLog.usuario_id).filter(
            AuditLog.tabla == AtencionDiaDetalle.__tablename__, AuditLog.accion == 'CORRECCION_EXCEL')
        if not ca._is_admin_user():
            clientes = ca._clientes_visibles_query(scope).with_entities(ca.ClienteComercial.id)
            query = query.filter(or_(
                AuditLog.datos_anteriores['cliente_id'].as_integer().in_(clientes),
                (AuditLog.datos_anteriores['cliente_id'].as_integer().is_(None)) &
                (AuditLog.datos_anteriores['vendedor_id'].as_integer() == scope.id)))
        for parametro, final in [('desde', False), ('hasta', True)]:
            if request.args.get(parametro):
                fecha = datetime.strptime(request.args[parametro], '%Y-%m-%d')
                if final:
                    fecha = fecha.replace(hour=23, minute=59, second=59, microsecond=999999)
                query = query.filter(AuditLog.created_at <= fecha if final else AuditLog.created_at >= fecha)
        if request.args.get('cliente_id'):
            query = query.filter(AuditLog.datos_anteriores['cliente_id'].as_integer() == int(request.args['cliente_id']))
        pagina = max(1, int(request.args.get('pagina', 1)))
        total = query.count()
        exportar = request.args.get('formato') == 'xlsx'
        query = query.order_by(AuditLog.created_at.desc(), AuditLog.id.desc())
        if not exportar:
            query = query.offset((pagina - 1) * 50).limit(50)
        filas = []
        for log, reg, usuario in query.all():
            anterior, nuevo = log.datos_anteriores or {}, log.datos_nuevos or {}
            meta = nuevo.get('_correccion', {})
            for campo in CAMPOS:
                if anterior.get(campo) == nuevo.get(campo):
                    continue
                filas.append({'fecha': log.created_at.isoformat() + 'Z', 'usuario': meta.get('usuario') or (usuario.usuario if usuario else str(log.usuario_id)),
                              'empresa': meta.get('empresa', ''), 'archivo': meta.get('archivo', ''),
                              'lote': meta.get('lote', ''), 'atencion_id': log.registro_id, 'campo': CAMPOS[campo],
                              'antes': anterior.get(campo), 'despues': nuevo.get(campo)})
        if exportar:
            wb = Workbook()
            ws = wb.active
            ws.title = 'Antes y despues'
            headers = ['fecha', 'usuario', 'empresa', 'archivo', 'lote', 'atencion_id', 'campo', 'antes', 'despues']
            ws.append(headers)
            for fila in filas:
                ws.append([fila[h] for h in headers])
                for cell in ws[ws.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            buf = BytesIO()
            wb.save(buf)
            buf.seek(0)
            return send_file(buf, as_attachment=True, download_name='Correcciones-atenciones.xlsx')
        return jsonify(filas=filas, total=total, pagina=pagina, paginas=(total + 49) // 50)
    except PermissionError as exc:
        return jsonify(error=str(exc)), 403
    except ValueError:
        return jsonify(error='Filtros invalidos'), 400
