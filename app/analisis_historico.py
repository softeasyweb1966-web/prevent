"""Análisis de solo lectura; el Excel no crea convenios ni paquetes.

Formato PREVENT_RELACIONES_HISTORICO v1:
_CONTROL contiene un manifiesto JSON fragmentado y su HMAC-SHA256, firmado
con la clave de la instancia exportadora. El futuro importador deberá verificar
firma, versión, cabeceras, número/orden de filas y hashes antes de aceptar nombres.
Solo PAQUETES.Nombre paquete queda fuera de los hashes y desbloqueado.
Después deberá revalidar catálogo, permisos y relaciones actuales y pedir
confirmación humana. No hay importador ni escritura de negocio en este módulo.
"""
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import hmac
import json
import uuid
from types import SimpleNamespace

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from sqlalchemy.orm import selectinload

from app.models import (AtencionDiaDetalle, ClienteComercialTarifa,
                        ComercialCatalogoItem)
from app.clientes_scope import clientes_visibles
from app.clientes_maestro import normalizar


FORMATO = 'PREVENT_RELACIONES_HISTORICO'
VERSION = 1


def _json(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'))


def _composicion(componentes):
    return tuple(sorted(componentes))


def _id_composicion(componentes):
    return 'COM-' + sha256(_json(componentes).encode('utf-8')).hexdigest()


def _valor_celda(value):
    if value is None:
        return ''
    if isinstance(value, (int, float, Decimal)):
        return format(Decimal(str(value)).normalize(), 'f')
    return str(value)


def manifiesto_hojas(wb):
    """Huella independiente de estilos; sobre valores, no sobre fórmulas evaluadas."""
    resultado = {}
    for ws in wb:
        if ws.title == '_CONTROL':
            continue
        headers = [_valor_celda(c.value) for c in ws[1]]
        editable = headers.index('Nombre paquete') if ws.title == 'PAQUETES' else -1
        hashes = []
        for row in ws.iter_rows(min_row=2):
            valores = [_valor_celda(c.value) if i != editable else '' for i, c in enumerate(row)]
            hashes.append(sha256(_json(valores).encode('utf-8')).hexdigest())
        resultado[ws.title] = {'columnas': headers, 'filas': hashes}
    return resultado


def _hoja(wb, nombre, columnas, filas):
    ws = wb.create_sheet(nombre)
    ws.append(columnas)
    for fila in filas:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in fila])
    for row in ws:
        for cell in row:
            # Los nombres procedentes del histórico nunca se ejecutan como fórmulas.
            if isinstance(cell.value, str):
                if len(cell.value) > 32767:
                    raise ValueError('El análisis excede el tamaño de una celda. Seleccione un cliente o un período menor.')
                cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='17365D')
    for index, nombre_col in enumerate(columnas, 1):
        letra = ws.cell(1, index).column_letter
        ws.column_dimensions[letra].width = min(55, max(19, len(nombre_col) + 3))
        if 'valor' in nombre_col.lower() or 'tarifa' in nombre_col.lower():
            for row in range(2, ws.max_row + 1):
                ws.cell(row, index).number_format = '#,##0.00'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    if nombre == 'PAQUETES':
        col = columnas.index('Nombre paquete') + 1
        for row in range(2, ws.max_row + 1):
            ws.cell(row, col).protection = Protection(locked=False)
            ws.cell(row, col).fill = PatternFill('solid', fgColor='FFF2CC')
    return ws


def generar_excel_historico(desde, hasta_exclusiva, cliente_id, vendedor, es_admin, usuario, clave):
    # Solo reutilizar funciones sin escrituras; NO llamar a generar_prefacturas().
    from app.routes.cargue_atenciones import (
        _construir_lookup_catalogo, _clasificar_servicio,
        _normalizar_forma_pago, _condicion_scope_atenciones)

    clientes = {c.id: c for c in clientes_visibles().all()}
    candidatos_clientes = defaultdict(set)
    for c in clientes.values():
        for nombre in (c.razon_social, c.nombre_comercial, c.nit, *(c.nombres_alternativos or [])):
            if normalizar(nombre):
                candidatos_clientes[normalizar(nombre)].add(c.id)
    if cliente_id is not None:
        if cliente_id not in clientes:
            raise PermissionError('No tiene acceso al cliente seleccionado.')
        clientes = {cliente_id: clientes[cliente_id]}
    items = {i.id: i for i in ComercialCatalogoItem.query.options(
        selectinload(ComercialCatalogoItem.paquete_componentes)).all()}
    por_nombre = defaultdict(list)
    paquetes = defaultdict(list)
    for item in items.values():
        por_nombre[item.nombre.upper().strip()].append(item)
        if item.tipo_item == 'PAQUETE' and item.paquete_componentes:
            paquetes[_composicion((p.examen_id, p.cantidad) for p in item.paquete_componentes)].append(item)
    tarifas = {(t.cliente_id, t.catalogo_item_id): t for t in ClienteComercialTarifa.query.filter(
        ClienteComercialTarifa.cliente_id.in_(clientes)).all()}
    query = AtencionDiaDetalle.query.filter(
        AtencionDiaDetalle.fecha_creacion_orden >= desde,
        AtencionDiaDetalle.fecha_creacion_orden < hasta_exclusiva)
    if not es_admin:
        query = query.filter(_condicion_scope_atenciones(vendedor))
    if cliente_id is not None:
        query = query.filter((AtencionDiaDetalle.cliente_id == cliente_id) | AtencionDiaDetalle.cliente_id.is_(None))
    registros = query.order_by(AtencionDiaDetalle.id).all()
    clasificaciones = _construir_lookup_catalogo()
    ordenes = defaultdict(list)
    incidencias, evidencia, origenes = [], [], {}

    def incidencia(r, motivo):
        cliente = clientes.get(r.cliente_id)
        incidencias.append([r.id, r.cliente_id, cliente.razon_social if cliente else
                            r.acuerdo_comercial or r.empresa_mision or 'Sin identificar',
                            r.servicio, motivo, r.cargue_id, r.archivo_origen])

    for r in registros:
        forma = _normalizar_forma_pago(r.forma_pago)
        bucket = 'CREDITO' if forma == 'CREDITO' else 'EFECTIVO' if forma in (
            'EFECTIVO', 'CONTADO', 'PARTICULAR', 'PARTICULARES') else None
        if (r.estado_orden or '').upper().strip() == 'ANULADA':
            incidencia(r, 'Excluida: orden anulada'); continue
        if _clasificar_servicio(r.servicio, clasificaciones) == 'ECOBABY' or 'ECOBABY' in (r.servicio or '').upper():
            incidencia(r, 'Excluida: ECOBABY (criterio de prefacturas)'); continue
        if bucket is None:
            incidencia(r, 'Excluida: forma de pago fuera del criterio de prefacturas'); continue
        cid = r.cliente_id
        if cid is None:
            nombre = (r.acuerdo_comercial or r.empresa_mision or '').strip()
            if not nombre:
                incidencia(r, 'Empresa sin identificar'); continue
            candidatos = set()
            for valor in (r.acuerdo_comercial, r.empresa_mision):
                candidatos.update(candidatos_clientes.get(normalizar(valor), set()))
            if cliente_id is not None and candidatos != {cliente_id}:
                continue
            # Agrupación de presentación, no asignación de cliente ni cambio del ORM.
            cid = 'ORIG-' + sha256(normalizar(nombre).encode('utf-8')).hexdigest()
            if cid not in clientes:
                clientes[cid] = SimpleNamespace(razon_social=nombre, nit='', examenes_convenidos='', servicios_convenidos='')
                origenes[cid] = set()
            origenes[cid].update(candidatos)
        elif cid not in clientes:
            incidencia(r, 'Empresa sin vínculo autorizado'); continue
        ordenes[(cid, bucket, r.nro_identificacion or '', r.nro_orden or f'SIN-{r.id}')].append(r)

    servicios = defaultdict(list)
    combinaciones = defaultdict(lambda: defaultdict(list))
    for orden, regs in ordenes.items():
        cid, bucket, paciente, numero = orden
        conflictos = []
        if not paciente or not regs[0].nro_orden:
            conflictos.append('Identificación de paciente u orden incompleta')
        if len({r.fecha_creacion_orden.date() for r in regs}) > 1 or len({r.sede for r in regs}) > 1:
            conflictos.append('Número de orden con fechas o sedes diferentes')
        if any(n > 1 for n in Counter((r.servicio or '').upper().strip() for r in regs).values()):
            conflictos.append('Examen repetido: revisar duplicado o cantidad real')
        componentes = Counter()
        resueltos = []
        for r in regs:
            candidatos = por_nombre.get((r.servicio or '').upper().strip(), [])
            item = candidatos[0] if len(candidatos) == 1 else None
            if item is None:
                conflictos.append('Servicio sin coincidencia única en catálogo')
            else:
                componentes[item.id] += 1
            resueltos.append((r, item))
        fecha = min(r.fecha_creacion_orden for r in regs).date()
        total = sum((r.precio for r in regs), Decimal('0')) if all(r.precio is not None for r in regs) else None
        observacion = {'orden': orden, 'paciente': paciente, 'fecha': fecha, 'valor': total,
                       'forma': bucket, 'ids': [r.id for r in regs]}
        for r, item in resueltos:
            evidencia.append([r.id, cid, clientes[cid].razon_social, r.nro_orden,
                              r.fecha_creacion_orden.isoformat(), r.servicio, item.id if item else None,
                              r.precio, bucket, r.cargue_id, r.archivo_origen])
            if item is not None and not conflictos:
                servicios[(cid, item.id)].append({**observacion, 'valor': r.precio, 'ids': [r.id]})
        if conflictos:
            for r in regs:
                incidencia(r, '; '.join(sorted(set(conflictos))))
            continue
        if len(regs) == 1 and resueltos[0][1].tipo_item == 'PAQUETE':
            item = resueltos[0][1]
            comp = _composicion((p.examen_id, p.cantidad) for p in item.paquete_componentes)
            if comp:
                combinaciones[comp][cid].append({**observacion, 'explicito': item.id})
            else:
                incidencia(regs[0], 'Paquete explícito sin composición en catálogo')
        elif len(componentes) >= 2:
            if all(items[i].tipo_item == 'EXAMEN' and items[i].clasificacion_completa for i in componentes):
                combinaciones[_composicion(componentes.items())][cid].append({**observacion, 'explicito': None})
            else:
                incidencia(regs[0], 'Combinación con ítems que no son exámenes completamente clasificados')

    def resumen(obs):
        meses = sorted({o['fecha'].strftime('%Y-%m') for o in obs})
        return [len({o['orden'] for o in obs}), ', '.join(meses), max(o['fecha'] for o in obs).isoformat(),
                len({(o['orden'][0], o['paciente']) for o in obs if o['paciente']})]

    def valores(obs):
        counts = Counter(o['valor'] for o in obs)
        return '; '.join(f'{v if v is not None else "Sin valor completo"}: {n} uso(s)'
                         for v, n in sorted(counts.items(), key=lambda p: (p[0] is None, p[0] or 0)))

    def alertas(obs):
        valores_validos = {o['valor'] for o in obs if o['valor'] is not None}
        mensajes = ['Valores históricos; no determinan tarifa vigente']
        if len(valores_validos) > 1:
            mensajes.append('Múltiples valores históricos')
        if any(o['valor'] is None or o['valor'] <= 0 for o in obs):
            mensajes.append('Precio nulo, cero o negativo: revisión obligatoria')
        return '; '.join(mensajes)

    filas_paquetes, filas_componentes, filas_empresas, filas_valores = [], [], [], []
    for n, (comp, empresas) in enumerate(sorted(combinaciones.items()), 1):
        ident, etiqueta = _id_composicion(comp), f'Paquete {n}'
        existentes = paquetes.get(comp, [])
        descripcion = ' + '.join(f'{items[i].nombre} × {cantidad}' for i, cantidad in comp)
        todas = [o for obs in empresas.values() for o in obs]
        filas_paquetes.append([ident, etiqueta, '', descripcion, len(empresas),
                               '; '.join(clientes[c].razon_social for c in sorted(empresas, key=str)),
                               *resumen(todas), valores(todas),
                               '; '.join(f'{p.id}: {p.nombre} ({"activo" if p.activo else "inactivo"})' for p in existentes),
                               'Coincidencia de composición: requiere confirmación' if existentes else 'Candidato; no creado'])
        for item_id, cantidad in comp:
            filas_componentes.append([ident, etiqueta, item_id, items[item_id].nombre, cantidad,
                                     items[item_id].tipo_examen, 'Activo' if items[item_id].activo else 'Inactivo'])
        for cid, obs in sorted(empresas.items(), key=lambda par: str(par[0])):
            recurrente = len({o['orden'] for o in obs}) >= 3 and len({o['fecha'].strftime('%Y-%m') for o in obs}) >= 2
            relaciones = [tarifas[(cid, p.id)] for p in existentes if (cid, p.id) in tarifas]
            relacion_txt = '; '.join(f'ID {t.id}, paquete {t.catalogo_item_id}, tarifa {t.tarifa_negociada}, '
                                     f'{"activa" if t.activo else "inactiva"}, '
                                     f'{t.vigencia_desde or "sin inicio"} / {t.vigencia_hasta or "sin fin"}' for t in relaciones)
            estado = 'Paquete explícito (composición actual)' if any(o['explicito'] for o in obs) else (
                'Paquete provisional recurrente' if recurrente else 'Combinación ocasional')
            notas = [alertas(obs)]
            if cid in origenes:
                notas.append('Empresa de origen sin vínculo; resolver identidad antes de crear relaciones')
            if len(existentes) > 1:
                notas.append('Varios paquetes existentes con la misma composición')
            if any(cantidad != 1 for _, cantidad in comp):
                notas.append('El formulario actual solo guarda cantidad 1')
            if any(not items[i].activo for i, _ in comp):
                notas.append('Componentes inactivos')
            if clientes[cid].examenes_convenidos or clientes[cid].servicios_convenidos:
                notas.append('Cliente con convenios anteriores: revisar antes de asignar')
            filas_empresas.append([ident, etiqueta, cid, clientes[cid].nit, clientes[cid].razon_social,
                                   estado, *resumen(obs), len({o['orden'] for o in obs}), valores(obs),
                                   relacion_txt, 'Pendiente identidad; verificar relaciones' if cid in origenes else
                                   'Ya existe; conservar' if relaciones else 'Propuesta nueva; no aplicada',
                                   '; '.join(notas)])
            for (forma, valor), subset in _por_valor(obs).items():
                filas_valores.append([ident, etiqueta, cid, clientes[cid].razon_social, forma, valor,
                                      len(subset), min(o['fecha'] for o in subset).isoformat(),
                                      max(o['fecha'] for o in subset).isoformat()])

    filas_servicios = []
    for (cid, item_id), obs in sorted(servicios.items(), key=lambda par: (str(par[0][0]), par[0][1])):
        item, tarifa = items[item_id], tarifas.get((cid, item_id))
        filas_servicios.append([cid, clientes[cid].nit, clientes[cid].razon_social, item_id, item.nombre,
                               item.tipo_item, item.tipo_examen, *resumen(obs), valores(obs),
                               tarifa.id if tarifa else None, tarifa.tarifa_negociada if tarifa else None,
                               str(tarifa.vigencia_desde or '') if tarifa else '',
                               str(tarifa.vigencia_hasta or '') if tarifa else '',
                               ('Activa' if tarifa.activo else 'Inactiva') if tarifa else
                               'Pendiente identidad; verificar relaciones' if cid in origenes else 'Sin relación',
                               alertas(obs) + ('; Empresa sin vínculo: resolver identidad' if cid in origenes else '')])
    wb = Workbook()
    wb.remove(wb.active)
    fecha_final = (hasta_exclusiva - timedelta(days=1)).date()
    _hoja(wb, 'LEEME', ['Concepto', 'Detalle'], [
        ['Formato', f'{FORMATO} v{VERSION}'], ['Período', f'{desde.date()} / {fecha_final}'],
        ['Alcance', 'Todos los clientes autorizados' if cliente_id is None else f'Cliente {cliente_id}'],
        ['Acción', 'Análisis sin escrituras de paquetes, relaciones ni atenciones.'],
        ['Edición permitida', 'Solo Nombre paquete en PAQUETES (celdas amarillas). No alterar composición, IDs, filas ni hojas.'],
        ['Nombre vacío', 'Propuesta pendiente de nombre; no significa crear automáticamente.'],
        ['Identidad', 'ID composición: SHA-256 de pares [ID examen, cantidad] ordenados. Paquete N es una etiqueta dentro de este archivo.'],
        ['Recurrencia', 'Criterio provisional: 3 órdenes o más en al menos 2 meses por empresa.'],
        ['Combinación', 'Se analiza la orden completa por empresa, paciente y forma de pago; no se infieren subconjuntos.'],
        ['Uso', 'Una orden-paciente por empresa y forma de pago. Pacientes se cuentan por empresa, no se deduplican entre empresas.'],
        ['Precios', 'Valores observados, no tarifas propuestas ni vigentes. Ver VALORES_PAQUETE para comparación por empresa y fecha.'],
        ['Paquetes explícitos', 'Su composición corresponde al catálogo actual; no acredita la composición histórica.'],
        ['Cobertura', f'{len(registros)} líneas leídas; {len(evidencia)} líneas elegibles; {len(incidencias)} incidencias/exclusiones.'],
        ['Exclusiones', 'Mismos criterios de pago, ANULADA y ECOBABY de prefacturas. Ver INCIDENCIAS; no se corrige el histórico.'],
        ['Empresas sin vínculo', 'ID ORIG identifica un texto de origen, no un cliente confirmado. Ver EMPRESAS_ORIGEN; las coincidencias son propuestas y no autorizan crear relaciones.'],
        ['Futura importación', 'No implementada en esta fase. Verificar firma, formato, composición y permisos; mostrar vista previa y pedir confirmación.'],
        ['Nombre ya existente', 'Mismo nombre y composición: proponer reutilizar. Nombre con otra composición: conflicto. Nunca sobrescribir relaciones.'],
        ['Firma', 'HMAC-SHA256 con clave de la instancia exportadora. Conservar _CONTROL y dicha clave para futura validación.'],
    ])
    _hoja(wb, 'PAQUETES', ['ID composición', 'Provisional', 'Nombre paquete', 'Componentes/exámenes',
                          'Cantidad empresas', 'Empresas', 'Usos/órdenes', 'Meses', 'Última utilización',
                          'Pacientes por empresa (suma)', 'Valores históricos', 'Paquetes existentes', 'Estado'], filas_paquetes)
    _hoja(wb, 'COMPONENTES', ['ID composición', 'Provisional', 'ID examen', 'Examen', 'Cantidad', 'Clasificación', 'Estado'], filas_componentes)
    _hoja(wb, 'EMPRESA_PAQUETE', ['ID composición', 'Provisional', 'ID empresa', 'NIT', 'Empresa', 'Clasificación',
                                 'Cantidad usos', 'Meses', 'Última utilización', 'Pacientes distintos', 'Órdenes distintas',
                                 'Valores históricos', 'Relaciones existentes', 'Propuesta', 'Observaciones'], filas_empresas)
    _hoja(wb, 'VALORES_PAQUETE', ['ID composición', 'Provisional', 'ID empresa', 'Empresa', 'Forma de pago',
                                 'Valor histórico', 'Cantidad usos', 'Primera fecha del valor', 'Última fecha del valor'], filas_valores)
    _hoja(wb, 'SERVICIOS', ['ID empresa', 'NIT', 'Empresa', 'ID ítem', 'Servicio/ítem', 'Tipo', 'Clasificación',
                           'Cantidad usos', 'Meses', 'Última utilización', 'Pacientes distintos', 'Valores históricos',
                           'ID relación', 'Tarifa configurada', 'Vigencia desde', 'Vigencia hasta', 'Estado relación', 'Observaciones'], filas_servicios)
    _hoja(wb, 'EVIDENCIA', ['ID atención', 'ID empresa', 'Empresa', 'Orden', 'Fecha orden', 'Servicio original',
                           'ID ítem', 'Precio histórico', 'Forma pago', 'ID cargue', 'Archivo origen'], evidencia)
    _hoja(wb, 'INCIDENCIAS', ['ID atención', 'ID empresa', 'Empresa/origen', 'Servicio', 'Motivo', 'ID cargue', 'Archivo origen'], incidencias)
    _hoja(wb, 'EMPRESAS_ORIGEN', ['ID empresa análisis', 'Empresa de origen', 'IDs clientes candidatos', 'Estado'],
          [[cid, clientes[cid].razon_social, ', '.join(str(i) for i in sorted(ids)),
            'Coincidencia única por confirmar' if len(ids) == 1 else 'Ambigua' if ids else 'Sin coincidencia']
           for cid, ids in sorted(origenes.items())])
    manifiesto = _json({'formato': FORMATO, 'version': VERSION, 'exportacion_id': str(uuid.uuid4()),
                       'generado_utc': datetime.now(timezone.utc).isoformat(),
                       'usuario_id': usuario.id, 'usuario': usuario.usuario, 'desde': str(desde.date()),
                       'hasta': str(fecha_final), 'cliente_id': cliente_id,
                       'hojas': manifiesto_hojas(wb)})
    firma = hmac.new(str(clave).encode('utf-8'), manifiesto.encode('utf-8'), 'sha256').hexdigest()
    control = wb.create_sheet('_CONTROL')
    control.append(['Formato', FORMATO]); control.append(['Versión', VERSION])
    control.append(['HMAC-SHA256', firma])
    for i in range(0, len(manifiesto), 30000):
        control.append(['Manifiesto JSON', manifiesto[i:i + 30000]])
    control.sheet_state = 'veryHidden'
    resultado = BytesIO()
    wb.save(resultado); resultado.seek(0)
    return resultado


def _por_valor(obs):
    grupos = defaultdict(list)
    for o in obs:
        grupos[(o['forma'], o['valor'])].append(o)
    return grupos
