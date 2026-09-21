"""Importación y consulta de la información contable exportada desde SIIGO."""

from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import re
import threading
import unicodedata
from zoneinfo import ZoneInfo

from flask import jsonify, request, send_file, current_app
from flask_login import current_user, login_required
from sqlalchemy import func, or_
from werkzeug.utils import secure_filename

# Estado en memoria de la importacion de clientes en segundo plano.
_ESTADO_CARGA_CLIENTES = {}
_CARGA_CLIENTES_LOCK = threading.Lock()

from app.models import (
    ClienteComercial,
    ChatCarteraHilo,
    ChatCarteraMensaje,
    SiigoCarga,
    SiigoCliente,
    SiigoComprobante,
    SiigoCuentaContable,
    SiigoCuentaReporte,
    SiigoMovimiento,
    SiigoSeguimientoCartera,
    SiigoComprobantePago,
    Vendedor,
    db,
)
from app.routes import contable_bp
from app.security import get_permission_names_for_user


from app.clientes_scope import es_administrador, clientes_visibles, vendedor_actual
from app.clientes_maestro import identificacion


def _clientes_informe():
    query = clientes_visibles()
    # El filtro nunca amplía el alcance del usuario vendedor.
    vendedor = request.args.get('vendedor_id')
    if vendedor == 'sin_asignar':
        query = query.filter(ClienteComercial.vendedor_id.is_(None))
    elif vendedor:
        query = query.filter(ClienteComercial.vendedor_id == int(vendedor))
    contacto = request.args.get('contacto_id', type=int)
    if contacto:
        from app.models import ContactoCliente
        query = query.filter(ClienteComercial.contactos.any(ContactoCliente.id == contacto))
    return query


def _alcance_movimiento():
    if es_administrador() and not request.args.get('vendedor_id') and not request.args.get('contacto_id'):
        return True
    nits = [identificacion(c.nit) for c in _clientes_informe() if c.nit]
    normalizado = func.upper(func.regexp_replace(func.split_part(SiigoMovimiento.identificacion, '-', 1), '[^a-zA-Z0-9]', '', 'g'))
    return normalizado.in_(nits)


def _comprobantes_visibles(query=None):
    query = SiigoComprobante.query if query is None else query
    return query.filter(SiigoComprobante.movimientos.any(_alcance_movimiento()))


@contable_bp.route('/filtros-clientes', methods=['GET'])
@login_required
def filtros_clientes():
    _requiere_ventas()
    from app.models import ContactoCliente, clientes_contactos
    visibles = clientes_visibles().with_entities(ClienteComercial.id).subquery()
    vendedores = db.session.query(Vendedor.id, Vendedor.nombre).join(
        ClienteComercial, ClienteComercial.vendedor_id == Vendedor.id,
    ).join(visibles, visibles.c.id == ClienteComercial.id).distinct().order_by(Vendedor.nombre.asc()).all()
    contactos = db.session.query(ContactoCliente.id, ContactoCliente.nombre, ContactoCliente.vendedor_id).join(
        clientes_contactos, clientes_contactos.c.contacto_id == ContactoCliente.id,
    ).join(visibles, visibles.c.id == clientes_contactos.c.cliente_id).distinct().order_by(ContactoCliente.nombre.asc()).all()
    return jsonify(es_administrador=es_administrador(),
                   vendedores=[{'id': pk, 'nombre': nombre} for pk, nombre in vendedores],
                   contactos=[{'id': pk, 'nombre': nombre, 'vendedor_id': vendedor_id} for pk, nombre, vendedor_id in contactos])


TIPOS_COMPROBANTE_PERMITIDOS = {'FV', 'RC', 'NC', 'ND', 'AC'}
CLASIFICACIONES_REPORTE_VENTAS = {'INGRESO', 'NOTA_CREDITO', 'IVA_GENERADO'}
REFERENCIA_FACTURA_RE = re.compile(r'\b(FV-\d+-[A-Z0-9]+(?:-[A-Z0-9]+)*)\b', re.IGNORECASE)
APLICACIONES_CARTERA = {'RC': 'recaudado', 'AC': 'ajustes_ac', 'NC': 'notas_credito', 'ND': 'notas_debito'}
FECHA_REFERENCIA_RE = re.compile(r'(?:fecha|date):\s*(\d{2}/\d{2}/\d{4})', re.IGNORECASE)


def _texto(value):
    return str(value).strip() if value is not None else ''


def _normalizar(value):
    text = unicodedata.normalize('NFD', _texto(value).lower())
    return ''.join(char for char in text if unicodedata.category(char) != 'Mn')


def _normalizar_encabezado(value):
    """Acepta encabezados SIIGO incluso cuando Excel reemplaza una tilde por U+FFFD."""
    encabezado = _normalizar(value).replace('\ufffd', '')
    equivalencias = {
        'fecha elaboracin': 'fecha elaboracion',
        'cdigo': 'codigo',
        'cdigo contable': 'codigo contable',
        'dbito': 'debito',
        'crdito': 'credito',
        'identificacin': 'identificacion',
        'descripcin': 'descripcion',
        'direccin': 'direccion',
        'categora': 'categoria',
        'relacin con': 'relacion con',
        'verificacin': 'verificacion',
    }
    return equivalencias.get(encabezado, encabezado)


def _decimal(value):
    if value in (None, ''):
        return Decimal('0')
    texto = str(value).strip()
    if ',' in texto:
        texto = texto.replace('.', '').replace(',', '.')
    try:
        return Decimal(texto)
    except InvalidOperation as exc:
        raise ValueError(f'Valor monetario inválido: {value}') from exc


def _fecha(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for formato in ('%d/%m/%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(_texto(value), formato).date()
        except ValueError:
            continue
    raise ValueError(f'Fecha inválida: {value}')


def _referencia_factura(descripcion):
    referencias = {match.upper() for match in REFERENCIA_FACTURA_RE.findall(_texto(descripcion))}
    return next(iter(referencias)) if len(referencias) == 1 else None


def _fecha_vencimiento(descripcion):
    match = FECHA_REFERENCIA_RE.search(_texto(descripcion))
    return _fecha(match.group(1)) if match else None


def _tipo_cartera(tipo, codigo):
    """SIIGO exporta los ajustes antiguos como CC-AC y los nuevos como AC-1."""
    return 'AC' if tipo == 'CC' and codigo == 'AC' else tipo


def _referencia_movimiento(comprobante, movimiento):
    campos = (
        (movimiento.descripcion, movimiento.detalle) if comprobante.tipo_documento == 'RC'
        else (movimiento.detalle, movimiento.descripcion)
    )
    for campo in campos:
        if REFERENCIA_FACTURA_RE.search(_texto(campo)):
            return _referencia_factura(campo)
    return None


def _es_base_factura(comprobante, referencia):
    return comprobante.tipo_documento == 'FV' and (
        not referencia or referencia == f'FV-{comprobante.codigo_comprobante}-{comprobante.numero_comprobante}'
    )


def _movimiento_cartera(comprobante, movimiento):
    return {
        'comprobante': f'{comprobante.tipo_documento}-{comprobante.codigo_comprobante}-{comprobante.numero_comprobante}',
        'secuencia': movimiento.secuencia,
        'fecha': comprobante.fecha_elaboracion.isoformat(),
        'debito': float(movimiento.debito),
        'credito': float(movimiento.credito),
    }


def _cancelaciones_cartera(fecha_corte, identificaciones=None):
    """Cruza toda línea de cartera con su factura, sin restringir el tipo de documento.

    Solo se toma la cuenta por cobrar: sumar también bancos, ingresos o retenciones
    duplicaría las contrapartidas del mismo asiento. La FV base se suma por separado.
    """
    query = db.session.query(SiigoComprobante, SiigoMovimiento).join(SiigoMovimiento).filter(
            _alcance_movimiento(),
        SiigoMovimiento.codigo_contable == '13050501',
        SiigoComprobante.fecha_elaboracion <= fecha_corte,
        SiigoMovimiento.credito != SiigoMovimiento.debito,
    )
    if identificaciones is not None:
        identificacion_normalizada = func.upper(func.regexp_replace(
            func.split_part(func.coalesce(SiigoMovimiento.identificacion, ''), '-', 1),
            r'[.\s]', '', 'g',
        ))
        query = query.filter(identificacion_normalizada.in_({_nit_cartera(valor) for valor in identificaciones}))
    for comprobante, movimiento in query.all():
        referencia = _referencia_movimiento(comprobante, movimiento)
        if _es_base_factura(comprobante, referencia):
            continue
        valor = movimiento.credito - movimiento.debito
        yield comprobante, movimiento, referencia, valor


def _mismo_tercero_cartera(factura, movimiento):
    return factura is not None and (
        not factura['identificacion'] or not movimiento.identificacion
        or _nit_cartera(factura['identificacion']) == _nit_cartera(movimiento.identificacion)
    )


def _estado_actualizacion_comprobantes():
    """Resume la vigencia con base en la fecha contable, no en la del archivo."""
    fecha_ultimo_comprobante = _comprobantes_visibles(db.session.query(func.max(SiigoComprobante.fecha_elaboracion))).scalar()
    fecha_minima_requerida = datetime.now(ZoneInfo('America/Bogota')).date() - timedelta(days=1)
    al_dia = bool(
        fecha_ultimo_comprobante
        and fecha_ultimo_comprobante >= fecha_minima_requerida
    )
    return {
        'fecha_ultimo_comprobante': fecha_ultimo_comprobante.isoformat() if fecha_ultimo_comprobante else None,
        'fecha_minima_requerida': fecha_minima_requerida.isoformat(),
        'al_dia': al_dia,
        'dias_atraso': max((fecha_minima_requerida - fecha_ultimo_comprobante).days, 0)
        if fecha_ultimo_comprobante else None,
    }


def _puede_usar_ventas():
    if getattr(current_user, 'is_easy', False) or getattr(getattr(current_user, 'role', None), 'nombre', None) == 'Administrador':
        return True
    permisos = get_permission_names_for_user(current_user)
    return bool({
        'menu_ventas',
        'menu_cartera',
        'menu_comercial',
        'comercial_section_cartera',
        'comercial_clientes_read',
    }.intersection(permisos))


def _requiere_ventas():
    if not _puede_usar_ventas():
        raise PermissionError('No tienes permiso para consultar o cargar información de ventas.')


def _leer_excel():
    archivo = request.files.get('archivo')
    if archivo is None or not archivo.filename:
        raise ValueError('Debe seleccionar un archivo Excel.')
    if not archivo.filename.lower().endswith('.xlsx'):
        raise ValueError('Solo se aceptan archivos .xlsx.')

    contenido = archivo.read()
    if not contenido:
        raise ValueError('El archivo está vacío.')
    try:
        import openpyxl
        libro = openpyxl.load_workbook(BytesIO(contenido), read_only=True, data_only=True)
        hoja = libro.active
        # Algunas exportaciones SIIGO declaran solo la columna A aunque tengan más datos.
        hoja.reset_dimensions()
        filas = list(hoja.iter_rows(values_only=True))
        libro.close()
    except Exception as exc:
        raise ValueError(f'No fue posible leer el archivo Excel: {exc}') from exc
    return archivo.filename, contenido, filas


def _indice_encabezados(filas, requeridos):
    required = {_normalizar_encabezado(value) for value in requeridos}
    for index, fila in enumerate(filas):
        headers = [_normalizar_encabezado(value) for value in fila]
        if required.issubset(set(headers)):
            return index, {header: position for position, header in enumerate(headers) if header}
    raise ValueError(f'No se encontraron las columnas requeridas: {", ".join(requeridos)}.')


def _valor(fila, columns, name):
    index = columns.get(_normalizar_encabezado(name))
    return fila[index] if index is not None and index < len(fila) else None


def _crear_carga(tipo_archivo, nombre_archivo, contenido):
    digest = sha256(contenido).hexdigest()
    existente = SiigoCarga.query.filter_by(hash_archivo=digest).first()
    if existente:
        if tipo_archivo == existente.tipo_archivo == 'COMPROBANTES':
            # Completa tipos antes omitidos (AC y CC-AC); conserva la clave de origen.
            return existente
        raise ValueError(f'Este archivo ya fue cargado el {existente.created_at:%Y-%m-%d %H:%M}.')
    carga = SiigoCarga(
        tipo_archivo=tipo_archivo,
        nombre_archivo=nombre_archivo,
        hash_archivo=digest,
        usuario_id=current_user.id,
    )
    db.session.add(carga)
    db.session.flush()
    return carga


def _extraer_comprobantes(filas, inicio):
    """Agrupa las secuencias de SIIGO bajo su fila separadora de comprobante."""
    actual = None
    for fila in filas[inicio:]:
        first = _texto(fila[0] if fila else None)
        if first.startswith('Comprobante:'):
            if actual and actual['lineas']:
                yield actual
            match = re.match(r'^Comprobante:\s*([^\-\s]+)-([^\-\s]+)-(.+?)\s*$', first)
            if not match:
                raise ValueError(f'Formato de comprobante no reconocido: {first}')
            tipo, codigo, numero = (group.strip() for group in match.groups())
            actual = {'tipo': tipo, 'codigo': codigo, 'numero': numero, 'lineas': []}
        elif actual and isinstance(fila[0] if fila else None, (int, float)):
            actual['lineas'].append(fila)
    if actual and actual['lineas']:
        yield actual


def _comprobante_relevante(documento, columns):
    return (
        _tipo_cartera(documento['tipo'], documento['codigo']) in TIPOS_COMPROBANTE_PERMITIDOS
        or any(_texto(_valor(linea, columns, 'Codigo contable')) == '13050501'
               for linea in documento['lineas'])
    )


def _guardar_comprobantes_en_lote(carga, filas, header_row, columns):
    """Evita miles de consultas individuales que pueden agotar el tiempo web."""
    documentos = list(_extraer_comprobantes(filas, header_row + 1))
    permitidos = [documento for documento in documentos if _comprobante_relevante(documento, columns)]
    existentes = {
        (tipo, codigo, numero)
        for tipo, codigo, numero in db.session.query(
            SiigoComprobante.tipo_documento,
            SiigoComprobante.codigo_comprobante,
            SiigoComprobante.numero_comprobante,
        ).filter(SiigoComprobante.tipo_documento.in_({documento['tipo'] for documento in permitidos})).all()
    }
    imported = movements = 0
    total_debito = total_credito = Decimal('0')

    for documento in permitidos:
        key = (documento['tipo'], documento['codigo'], documento['numero'])
        if key in existentes:
            continue
        fecha = _fecha(_valor(documento['lineas'][0], columns, 'Fecha elaboracion'))
        document_debito = sum((_decimal(_valor(line, columns, 'Debito')) for line in documento['lineas']), Decimal('0'))
        document_credito = sum((_decimal(_valor(line, columns, 'Credito')) for line in documento['lineas']), Decimal('0'))
        if document_debito.quantize(Decimal('0.01')) != document_credito.quantize(Decimal('0.01')):
            reference = '-'.join(key)
            raise ValueError(f'El comprobante {reference} no cuadra: dÃ©bito {document_debito} / crÃ©dito {document_credito}.')

        comprobante = SiigoComprobante(
            tipo_documento=documento['tipo'], codigo_comprobante=documento['codigo'], numero_comprobante=documento['numero'],
            fecha_elaboracion=fecha, total_debito=document_debito, total_credito=document_credito, carga_id=carga.id,
        )
        for line in documento['lineas']:
            comprobante.movimientos.append(SiigoMovimiento(
                secuencia=int(_valor(line, columns, 'Secuencia')),
                codigo_contable=_texto(_valor(line, columns, 'Codigo contable')),
                cuenta_contable=_texto(_valor(line, columns, 'Cuenta contable')),
                identificacion=_texto(_valor(line, columns, 'Identificacion')) or None,
                sucursal=_texto(_valor(line, columns, 'Sucursal')) or None,
                nombre_tercero=_texto(_valor(line, columns, 'Nombre tercero')) or None,
                descripcion=_texto(_valor(line, columns, 'Descripcion')) or None,
                detalle=_texto(_valor(line, columns, 'Detalle')) or None,
                centro_costo=_texto(_valor(line, columns, 'Centro de costo')) or None,
                debito=_decimal(_valor(line, columns, 'Debito')),
                credito=_decimal(_valor(line, columns, 'Credito')),
            ))
            movements += 1
        db.session.add(comprobante)
        existentes.add(key)
        imported += 1
        total_debito += document_debito
        total_credito += document_credito

    if not imported:
        raise ValueError('El archivo no contiene comprobantes nuevos de ventas o con movimientos de cartera. No se duplicaron registros.')
    carga.registros_leidos = len(documentos)
    carga.registros_importados += imported
    carga.registros_omitidos = len(documentos) - carga.registros_importados
    carga.total_debito += total_debito
    carga.total_credito += total_credito
    db.session.commit()
    return jsonify({'mensaje': 'Comprobantes cargados correctamente.', 'comprobantes': imported, 'movimientos': movements, 'omitidos': len(documentos) - imported})


@contable_bp.route('/resumen', methods=['GET'])
@login_required
def resumen():
    try:
        _requiere_ventas()
        return jsonify({
            'clientes': _clientes_informe().count(),
            'clientes_sin_vendedor': _clientes_informe().filter(ClienteComercial.vendedor_id.is_(None)).count(),
            'cuentas': SiigoCuentaContable.query.count(),
            'comprobantes': _comprobantes_visibles().count(),
            'movimientos': SiigoMovimiento.query.filter(_alcance_movimiento()).count(),
            'vigencia_comprobantes': _estado_actualizacion_comprobantes(),
            'cargas': [
                {
                    'tipo': carga.tipo_archivo,
                    'archivo': carga.nombre_archivo,
                    'fecha': carga.created_at.strftime('%Y-%m-%d %H:%M'),
                    'importados': carga.registros_importados,
                    'omitidos': carga.registros_omitidos,
                }
                for carga in (SiigoCarga.query.order_by(SiigoCarga.created_at.desc(), SiigoCarga.id.desc()).limit(10).all() if es_administrador() else [])
            ],
        })
    except PermissionError as exc:
        return jsonify({'error': str(exc)}), 403


@contable_bp.route('/cargar-clientes', methods=['POST'])
@login_required
def cargar_clientes():
    """Importa el maestro de clientes desde el Excel EN SEGUNDO PLANO.

    El import es largo (miles de clientes/contactos); ejecutarlo dentro del
    request agotaria el timeout HTTP. Se lanza en un hilo con su propio contexto
    de aplicacion y se responde de inmediato; el avance se consulta con
    GET /estado-carga-clientes."""
    try:
        _requiere_ventas()
        from app.routes.comercial import _is_admin_user
        if not _is_admin_user():
            raise PermissionError('La importacion global del maestro requiere un administrador.')
        if str(request.form.get('reemplazar', '')).strip().lower() in {'1', 'true', 'si', 'yes'}:
            raise ValueError('El maestro se actualiza conservando clientes y relaciones. El reemplazo destructivo no está permitido.')

        with _CARGA_CLIENTES_LOCK:
            if _ESTADO_CARGA_CLIENTES.get('estado') == 'en_proceso':
                return jsonify({'error': 'Ya hay una importacion de clientes en curso. Espera a que termine.'}), 409

        # Leer el archivo AQUI (mientras el request esta vivo); el hilo no tiene request.
        nombre_archivo, contenido, filas = _leer_excel()
        usuario_id = current_user.id
        app_obj = current_app._get_current_object()

        with _CARGA_CLIENTES_LOCK:
            _ESTADO_CARGA_CLIENTES.clear()
            _ESTADO_CARGA_CLIENTES.update({
                'estado': 'en_proceso',
                'archivo': nombre_archivo,
                'iniciado': datetime.now(ZoneInfo('America/Bogota')).isoformat(),
                'resumen': None,
                'error': None,
            })

        hilo = threading.Thread(
            target=_ejecutar_carga_clientes,
            args=(app_obj, filas, contenido, nombre_archivo, usuario_id),
            daemon=True,
        )
        hilo.start()
        return jsonify({'mensaje': 'Importacion de clientes iniciada. Consulta el estado para ver el avance.', 'estado': 'en_proceso'}), 202
    except (ValueError, PermissionError) as exc:
        return jsonify({'error': str(exc)}), 400 if isinstance(exc, ValueError) else 403
    except Exception as exc:
        return jsonify({'error': f'No fue posible iniciar la carga de clientes: {exc}'}), 500


def _ejecutar_carga_clientes(app_obj, filas, contenido, nombre_archivo, usuario_id):
    from app.clientes_maestro import importar_clientes
    with app_obj.app_context():
        try:
            resumen = importar_clientes(filas, contenido, nombre_archivo, usuario_id)
            db.session.commit()
            with _CARGA_CLIENTES_LOCK:
                _ESTADO_CARGA_CLIENTES.update({
                    'estado': 'completado',
                    'resumen': resumen,
                    'finalizado': datetime.now(ZoneInfo('America/Bogota')).isoformat(),
                })
        except Exception as exc:
            db.session.rollback()
            with _CARGA_CLIENTES_LOCK:
                _ESTADO_CARGA_CLIENTES.update({
                    'estado': 'error',
                    'error': str(exc),
                    'finalizado': datetime.now(ZoneInfo('America/Bogota')).isoformat(),
                })
        finally:
            db.session.remove()


@contable_bp.route('/estado-carga-clientes', methods=['GET'])
@login_required
def estado_carga_clientes():
    try:
        _requiere_ventas()
        from app.routes.comercial import _is_admin_user
        if not _is_admin_user():
            raise PermissionError('Solo un administrador puede consultar la carga del maestro.')
        with _CARGA_CLIENTES_LOCK:
            estado = dict(_ESTADO_CARGA_CLIENTES) if _ESTADO_CARGA_CLIENTES else {'estado': 'inactivo'}
        return jsonify(estado)
    except PermissionError as exc:
        return jsonify({'error': str(exc)}), 403


@contable_bp.route('/cargar-cuentas', methods=['POST'])
@login_required
def cargar_cuentas():
    try:
        _requiere_ventas()
        nombre_archivo, contenido, filas = _leer_excel()
        header_row, columns = _indice_encabezados(filas, ['Código', 'Nombre'])
        carga = _crear_carga('CUENTAS', nombre_archivo, contenido)
        creados = actualizados = 0

        for fila in filas[header_row + 1:]:
            codigo = _texto(_valor(fila, columns, 'Código'))
            nombre = _texto(_valor(fila, columns, 'Nombre'))
            if not codigo or not nombre:
                continue
            cuenta = SiigoCuentaContable.query.filter_by(codigo=codigo).first()
            fields = {
                'nombre': nombre,
                'categoria': _texto(_valor(fila, columns, 'Categoría')) or None,
                'clase': _texto(_valor(fila, columns, 'Clase')) or None,
                'relacion_con': _texto(_valor(fila, columns, 'Relación con')) or None,
                'maneja_vencimientos': _texto(_valor(fila, columns, 'Maneja vencimientos')) or None,
                'activo': _normalizar(_valor(fila, columns, 'Activo')) in {'si', 'sí', 'true', '1'},
                'carga_id': carga.id,
            }
            if cuenta is None:
                db.session.add(SiigoCuentaContable(codigo=codigo, **fields))
                creados += 1
            else:
                for field, value in fields.items():
                    setattr(cuenta, field, value)
                actualizados += 1

        carga.registros_leidos = len(filas) - header_row - 1
        carga.registros_importados = creados + actualizados
        db.session.commit()
        return jsonify({'mensaje': 'Plan de cuentas cargado correctamente.', 'creados': creados, 'actualizados': actualizados})
    except (ValueError, PermissionError) as exc:
        db.session.rollback()
        return jsonify({'error': str(exc)}), 400 if isinstance(exc, ValueError) else 403
    except Exception as exc:
        db.session.rollback()
        return jsonify({'error': f'No fue posible cargar cuentas: {exc}'}), 500


@contable_bp.route('/cargar-comprobantes', methods=['POST'])
@login_required
def cargar_comprobantes():
    try:
        _requiere_ventas()
        nombre_archivo, contenido, filas = _leer_excel()
        try:
            header_row, columns = _indice_encabezados(filas, ['Secuencia', 'Fecha elaboración', 'Código contable', 'Débito', 'Crédito'])
        except ValueError as exc:
            raise ValueError(
                'No se reconoce el formato del archivo. Para cargar comprobantes, '
                'en SIIGO vaya a Reportes - Versiones anteriores reportes - Comprobantes detallados, use Agrupar, '
                'seleccione el período que desea actualizar y exporte el informe a Excel (.xlsx). '
                'El informe Consecutivo de comprobantes no contiene el detalle requerido. '
                f'Detalle de validación: {exc}'
            ) from exc
        carga = _crear_carga('COMPROBANTES', nombre_archivo, contenido)
        return _guardar_comprobantes_en_lote(carga, filas, header_row, columns)

    except (ValueError, PermissionError) as exc:
        db.session.rollback()
        return jsonify({'error': str(exc)}), 400 if isinstance(exc, ValueError) else 403
    except Exception as exc:
        db.session.rollback()
        return jsonify({'error': f'No fue posible cargar comprobantes: {exc}'}), 500


@contable_bp.route('/clientes', methods=['GET'])
@login_required
def consultar_clientes():
    try:
        _requiere_ventas()
        search = _texto(request.args.get('q'))
        query = _clientes_informe()
        if search:
            like = f'%{search}%'
            query = query.filter(or_(SiigoCliente.identificacion.ilike(like), SiigoCliente.nombre.ilike(like)))
        items = query.order_by(SiigoCliente.nombre).all()
        clientes = [{'id': item.id, 'identificacion': item.identificacion, 'sucursal': item.sucursal, 'nombre': item.nombre, 'ciudad': item.ciudad, 'estado': item.estado,
                     'vendedor_id': item.vendedor_id, 'vendedor_nombre': item.vendedor.nombre if item.vendedor else 'Sin vendedor asignado'} for item in items]
        return jsonify({'clientes': clientes})
    except PermissionError as exc:
        return jsonify({'error': str(exc)}), 403


@contable_bp.route('/comprobantes', methods=['GET'])
@login_required
def consultar_comprobantes():
    try:
        _requiere_ventas()
        query = _comprobantes_visibles()
        tipo = _texto(request.args.get('tipo')).upper()
        numero = _texto(request.args.get('numero'))
        desde = _texto(request.args.get('desde'))
        hasta = _texto(request.args.get('hasta'))
        cliente = _texto(request.args.get('cliente'))
        if tipo == 'AC':
            query = query.filter(or_(
                SiigoComprobante.tipo_documento == 'AC',
                (SiigoComprobante.tipo_documento == 'CC') & (SiigoComprobante.codigo_comprobante == 'AC'),
            ))
        elif tipo:
            query = query.filter_by(tipo_documento=tipo)
        if numero:
            query = query.filter(SiigoComprobante.numero_comprobante.ilike(f'%{numero}%'))
        if desde:
            query = query.filter(SiigoComprobante.fecha_elaboracion >= _fecha(desde))
        if hasta:
            query = query.filter(SiigoComprobante.fecha_elaboracion <= _fecha(hasta))
        if cliente:
            like = f'%{cliente}%'
            query = query.join(SiigoMovimiento).filter(_alcance_movimiento(), or_(SiigoMovimiento.identificacion.ilike(like), SiigoMovimiento.nombre_tercero.ilike(like))).distinct()
        items = query.order_by(SiigoComprobante.fecha_elaboracion.desc()).limit(200).all()
        movimientos_visibles = {item.id: SiigoMovimiento.query.filter(SiigoMovimiento.comprobante_id == item.id, _alcance_movimiento()).all() for item in items}
        return jsonify({'comprobantes': [{
            'tipo': item.tipo_documento, 'codigo': item.codigo_comprobante, 'numero': item.numero_comprobante,
            'fecha': item.fecha_elaboracion.isoformat(), 'debito': float(sum(m.debito for m in movimientos_visibles[item.id])), 'credito': float(sum(m.credito for m in movimientos_visibles[item.id])),
            'movimientos': len(movimientos_visibles[item.id]),
        } for item in items]})
    except (ValueError, PermissionError) as exc:
        return jsonify({'error': str(exc)}), 400 if isinstance(exc, ValueError) else 403

@contable_bp.route('/comparativo-clientes', methods=['GET'])
@login_required
def comparativo_clientes():
    try:
        _requiere_ventas()
        periodo_a_desde = _fecha(request.args.get('periodo_a_desde'))
        periodo_a_hasta = _fecha(request.args.get('periodo_a_hasta'))
        periodo_b_desde = _fecha(request.args.get('periodo_b_desde'))
        periodo_b_hasta = _fecha(request.args.get('periodo_b_hasta'))
        fecha_corte_cartera = _fecha(request.args.get('fecha_corte_cartera')) if request.args.get('fecha_corte_cartera') else periodo_b_hasta
        if periodo_a_desde > periodo_a_hasta or periodo_b_desde > periodo_b_hasta:
            raise ValueError('La fecha inicial no puede ser posterior a la fecha final.')

        def terceros_del_periodo(desde, hasta):
            rows = db.session.query(
                SiigoMovimiento.identificacion,
                func.max(SiigoMovimiento.nombre_tercero),
                func.count(func.distinct(SiigoComprobante.id)),
                func.coalesce(func.sum(SiigoMovimiento.debito - SiigoMovimiento.credito), 0),
            ).join(SiigoComprobante).filter(
            _alcance_movimiento(),
                SiigoComprobante.tipo_documento == 'FV',
                SiigoComprobante.fecha_elaboracion.between(desde, hasta),
                SiigoMovimiento.identificacion.isnot(None),
                SiigoMovimiento.codigo_contable == '13050501',
            ).group_by(SiigoMovimiento.identificacion).all()
            return {
                identificacion: {
                    'identificacion': identificacion,
                    'nombre': nombre or '',
                    'facturas': int(facturas),
                    'facturacion': valor or Decimal('0'),
                }
                for identificacion, nombre, facturas, valor in rows
            }

        periodo_a = terceros_del_periodo(periodo_a_desde, periodo_a_hasta)
        periodo_b = terceros_del_periodo(periodo_b_desde, periodo_b_hasta)
        nuevos = sorted((periodo_b[key] for key in periodo_b.keys() - periodo_a.keys()), key=lambda item: item['nombre'])
        no_volvieron = sorted((periodo_a[key] for key in periodo_a.keys() - periodo_b.keys()), key=lambda item: item['nombre'])

        identificaciones = list((periodo_b.keys() - periodo_a.keys()) | (periodo_a.keys() - periodo_b.keys()))
        facturas_cartera = {}
        pagos_sin_factura = {}
        if identificaciones:
            filas_factura = db.session.query(SiigoComprobante, SiigoMovimiento).join(SiigoMovimiento).filter(
            _alcance_movimiento(),
                SiigoComprobante.tipo_documento == 'FV',
                SiigoMovimiento.codigo_contable == '13050501',
                SiigoMovimiento.identificacion.in_(identificaciones),
                SiigoComprobante.fecha_elaboracion <= fecha_corte_cartera,
                SiigoMovimiento.debito != SiigoMovimiento.credito,
            ).all()
            for comprobante, movimiento in filas_factura:
                referencia = _referencia_movimiento(comprobante, movimiento)
                if not _es_base_factura(comprobante, referencia):
                    continue
                referencia = referencia or f'FV-{comprobante.codigo_comprobante}-{comprobante.numero_comprobante}'
                factura = facturas_cartera.setdefault(referencia, {
                    'identificacion': movimiento.identificacion,
                    'valor': Decimal('0'),
                    'pagado': Decimal('0'),
                })
                factura['valor'] += movimiento.debito - movimiento.credito

            for comprobante, movimiento, referencia, valor in _cancelaciones_cartera(fecha_corte_cartera, identificaciones):
                factura = facturas_cartera.get(referencia)
                if not _mismo_tercero_cartera(factura, movimiento):
                    pagos_sin_factura[movimiento.identificacion] = pagos_sin_factura.get(movimiento.identificacion, Decimal('0')) + valor
                else:
                    factura['pagado'] += valor

        def totales(items):
            identificaciones_grupo = {item['identificacion'] for item in items}
            cartera = Decimal('0')
            pagos_pendientes = sum((pagos_sin_factura.get(identificacion, Decimal('0')) for identificacion in identificaciones_grupo), Decimal('0'))
            for factura in facturas_cartera.values():
                if factura['identificacion'] not in identificaciones_grupo:
                    continue
                saldo = factura['valor'] - factura['pagado']
                cartera += max(saldo, Decimal('0'))
                pagos_pendientes += max(-saldo, Decimal('0'))
            return {
                'facturacion': float(sum((item['facturacion'] for item in items), Decimal('0'))),
                'cartera': float(cartera),
                'pagos_pendientes_conciliar': float(pagos_pendientes),
            }

        return jsonify({
            'clientes_periodo_a': len(periodo_a),
            'clientes_periodo_b': len(periodo_b),
            'nuevos': [{**item, 'facturacion': float(item['facturacion'])} for item in nuevos],
            'no_volvieron': [{**item, 'facturacion': float(item['facturacion'])} for item in no_volvieron],
            'totales_nuevos': totales(nuevos),
            'totales_no_volvieron': totales(no_volvieron),
            'fecha_cartera': fecha_corte_cartera.isoformat(),
        })
    except (ValueError, PermissionError) as exc:
        return jsonify({'error': str(exc)}), 400 if isinstance(exc, ValueError) else 403


@contable_bp.route('/configuracion-ventas', methods=['GET', 'POST'])
@login_required
def configuracion_ventas():
    try:
        _requiere_ventas()
        if request.method == 'POST':
            data = request.get_json() or {}
            codigo = _texto(data.get('codigo_contable'))
            clasificacion = _texto(data.get('clasificacion')).upper()
            if not codigo or clasificacion not in CLASIFICACIONES_REPORTE_VENTAS:
                raise ValueError('Debe indicar una cuenta y una clasificación válida.')
            item = SiigoCuentaReporte.query.filter_by(codigo_contable=codigo).first()
            if item is None:
                item = SiigoCuentaReporte(codigo_contable=codigo, clasificacion=clasificacion)
                db.session.add(item)
            item.clasificacion = clasificacion
            item.activo = bool(data.get('activo', True))
            db.session.commit()

        configuradas = SiigoCuentaReporte.query.order_by(SiigoCuentaReporte.clasificacion, SiigoCuentaReporte.codigo_contable).all()
        nombres = {
            cuenta.codigo: cuenta.nombre
            for cuenta in SiigoCuentaContable.query.filter(
                SiigoCuentaContable.codigo.in_([item.codigo_contable for item in configuradas])
            ).all()
        } if configuradas else {}
        return jsonify({'cuentas': [{
            'codigo': item.codigo_contable,
            'nombre': nombres.get(item.codigo_contable, ''),
            'clasificacion': item.clasificacion,
            'activo': item.activo,
        } for item in configuradas]})
    except (ValueError, PermissionError) as exc:
        db.session.rollback()
        return jsonify({'error': str(exc)}), 400 if isinstance(exc, ValueError) else 403


@contable_bp.route('/ventas-mensuales', methods=['GET'])
@login_required
def ventas_mensuales():
    try:
        _requiere_ventas()
        anio = int(request.args.get('anio', datetime.now().year))
        if anio < 2000 or anio > 2100:
            raise ValueError('El año debe estar entre 2000 y 2100.')
        incluir_nc = _texto(request.args.get('incluir_nc')).lower() in {'1', 'true', 'si', 'yes', 'on'}
        incluir_iva = _texto(request.args.get('incluir_iva')).lower() in {'1', 'true', 'si', 'yes', 'on'}
        configuracion = SiigoCuentaReporte.query.filter_by(activo=True).all()
        ingresos = {item.codigo_contable for item in configuracion if item.clasificacion == 'INGRESO'}
        notas_credito = {item.codigo_contable for item in configuracion if item.clasificacion == 'NOTA_CREDITO'}
        iva = {item.codigo_contable for item in configuracion if item.clasificacion == 'IVA_GENERADO'}
        if not ingresos:
            raise ValueError('No hay cuentas de ingreso configuradas para el informe.')
        cuentas = set(ingresos)
        tipos = {'FV'}
        if incluir_nc:
            cuentas.update(notas_credito)
            tipos.add('NC')
        if incluir_iva:
            cuentas.update(iva)

        valor = func.coalesce(SiigoMovimiento.credito, 0) - func.coalesce(SiigoMovimiento.debito, 0)
        mes = func.extract('month', SiigoComprobante.fecha_elaboracion)
        rows = db.session.query(mes, func.sum(valor)).join(SiigoComprobante).filter(
            _alcance_movimiento(),
            SiigoComprobante.tipo_documento.in_(tipos),
            func.extract('year', SiigoComprobante.fecha_elaboracion) == anio,
            SiigoMovimiento.codigo_contable.in_(cuentas),
        ).group_by(mes).order_by(mes).all()
        totales = {int(numero_mes): float(total or 0) for numero_mes, total in rows}
        return jsonify({
            'anio': anio,
            'incluir_nc': incluir_nc,
            'incluir_iva': incluir_iva,
            'cuentas': sorted(cuentas),
            'meses': [{'mes': mes_numero, 'valor': totales.get(mes_numero, 0)} for mes_numero in range(1, 13)],
            'total': sum(totales.values()),
        })
    except (ValueError, PermissionError) as exc:
        return jsonify({'error': str(exc)}), 400 if isinstance(exc, ValueError) else 403


def _nit_cartera(value):
    return re.sub(r'[.\s]', '', _texto(value).split('-', 1)[0]).upper()


def _estado_compromiso(registro, hoy=None):
    if not registro.fecha_compromiso:
        return 'sin_compromiso'
    if registro.compromiso_cumplido_at:
        return 'cumplido'
    dias = (registro.fecha_compromiso - (hoy or datetime.now(ZoneInfo('America/Bogota')).date())).days
    return 'vencido' if dias < 0 else 'proximo' if dias <= 3 else 'pendiente'


ESTADOS_GESTION_CARTERA = {'SIN_GESTION', 'NO_LOCALIZADO', 'EN_PROCESO', 'CON_COMPROMISO'}
TIPOS_CHAT_CARTERA = {'AUTORIZACION_ARREGLO', 'SOLICITUD_INFO', 'GENERAL'}
ESTADOS_CHAT_CARTERA = {'ABIERTO', 'AUTORIZADO', 'RECHAZADO', 'CERRADO'}


def _serializar_seguimiento_cartera(registro):
    return {
        'id': registro.id,
        'identificacion': registro.identificacion,
        'estado_gestion': registro.estado_gestion,
        'comprobantes_pago': [{'id': a.id, 'nombre': a.nombre, 'tamano_bytes': a.tamano_bytes,
                              'registrado_por': a.usuario_nombre,
                              'url': f'/api/contable/seguimiento-cartera/comprobantes/{a.id}'} for a in registro.comprobantes_pago],
        'estado_compromiso': _estado_compromiso(registro),
        'compromiso_cumplido_at': registro.compromiso_cumplido_at.isoformat() + 'Z' if registro.compromiso_cumplido_at else None,
        'compromiso_cumplido_por': registro.compromiso_cumplido_por,
        'fecha_gestion': registro.fecha_gestion.isoformat(),
        'medio': registro.medio,
        'contacto': registro.contacto,
        'observaciones': registro.observaciones,
        'fecha_compromiso': registro.fecha_compromiso.isoformat() if registro.fecha_compromiso else None,
        'valor_compromiso': float(registro.valor_compromiso) if registro.valor_compromiso is not None else None,
        'proximo_seguimiento': registro.proximo_seguimiento.isoformat() if registro.proximo_seguimiento else None,
        'registrado_por': registro.usuario_nombre,
        'created_at': registro.created_at.isoformat() + 'Z',
    }


def _cliente_maestro_por_identificacion(identificacion_raw):
    clave = _nit_cartera(identificacion_raw)
    if not clave:
        return None
    for cliente in clientes_visibles().all():
        if _nit_cartera(cliente.nit) == clave:
            return cliente
    return None


def _vendedor_chat_visible(vendedor_id=None, cliente=None):
    if cliente is not None:
        vendedor_id = cliente.vendedor_id
    if vendedor_id in ('', None):
        vendedor_id = None
    if vendedor_id is None:
        vendedor = vendedor_actual()
        if vendedor:
            return vendedor
        raise PermissionError('Seleccione un vendedor para la conversacion.')
    vendedor = Vendedor.query.get(int(vendedor_id))
    if vendedor is None:
        raise ValueError('Seleccione un vendedor valido.')
    if not es_administrador():
        actual = vendedor_actual()
        if actual is None or actual.id != vendedor.id:
            raise PermissionError('No tienes acceso a este vendedor.')
    return vendedor


def _hilo_chat_visible(hilo_id):
    hilo = ChatCarteraHilo.query.get(hilo_id)
    if hilo is None:
        return None
    if es_administrador():
        return hilo
    actual = vendedor_actual()
    if actual is None or hilo.vendedor_id != actual.id:
        return None
    return hilo


def _serializar_chat_cartera(hilo, incluir_mensajes=True):
    mensajes = sorted(hilo.mensajes, key=lambda item: (item.created_at, item.id)) if incluir_mensajes else []
    return {
        'id': hilo.id,
        'tipo': hilo.tipo,
        'estado': hilo.estado,
        'asunto': hilo.asunto,
        'cliente_id': hilo.cliente_id,
        'identificacion': hilo.identificacion,
        'cliente_nombre': hilo.cliente_nombre,
        'vendedor_id': hilo.vendedor_id,
        'vendedor_nombre': hilo.vendedor.nombre if hilo.vendedor else '',
        'creado_por': hilo.creado_por.nombre_completo if hilo.creado_por else '',
        'autorizado_por': hilo.autorizado_por.nombre_completo if hilo.autorizado_por else None,
        'autorizado_at': hilo.autorizado_at.isoformat() + 'Z' if hilo.autorizado_at else None,
        'cerrado_at': hilo.cerrado_at.isoformat() + 'Z' if hilo.cerrado_at else None,
        'created_at': hilo.created_at.isoformat() + 'Z',
        'updated_at': hilo.updated_at.isoformat() + 'Z',
        'mensajes': [{
            'id': mensaje.id,
            'remitente_id': mensaje.remitente_id,
            'remitente_nombre': mensaje.remitente_nombre,
            'mensaje': mensaje.mensaje,
            'decision': mensaje.decision,
            'created_at': mensaje.created_at.isoformat() + 'Z',
        } for mensaje in mensajes],
    }


def _clientes_agrupados_por_responsable(clientes):
    grupos = {}
    for cliente in clientes:
        clave = _nit_cartera(cliente.nit)
        responsable = _texto(cliente.responsable)
        if not clave or not responsable:
            continue
        grupo_clave = _normalizar(responsable)
        grupos.setdefault(grupo_clave, {
            'responsable': responsable,
            'telefono_responsable': cliente.telefono_responsable,
            'empresas': [],
        })['empresas'].append({
            'identificacion': clave,
            'cliente': cliente.razon_social,
            'vendedor': cliente.vendedor.nombre if cliente.vendedor else 'Sin asignar',
        })
    resultado = {}
    for grupo in grupos.values():
        grupo['empresas'].sort(key=lambda item: (_normalizar(item['cliente']), item['identificacion']))
        for empresa in grupo['empresas']:
            resultado[empresa['identificacion']] = grupo
    return resultado


@contable_bp.route('/chat-cartera', methods=['GET', 'POST'])
@login_required
def chat_cartera():
    try:
        _requiere_ventas()
        if request.method == 'GET':
            identificacion_raw = _texto(request.args.get('identificacion'))
            vendedor_id = request.args.get('vendedor_id', type=int)
            query = ChatCarteraHilo.query
            if identificacion_raw:
                clave = _nit_cartera(identificacion_raw)
                cliente = _cliente_maestro_por_identificacion(clave)
                if cliente is None:
                    # Puede existir cartera SIIGO sin ficha completa; el alcance se valida con movimientos visibles.
                    identificacion_normalizada = func.upper(func.regexp_replace(
                        func.split_part(SiigoMovimiento.identificacion, '-', 1),
                        '[^a-zA-Z0-9]',
                        '',
                        'g',
                    ))
                    existe = db.session.query(SiigoMovimiento.id).join(SiigoComprobante).filter(
                        _alcance_movimiento(),
                        identificacion_normalizada == clave,
                        SiigoComprobante.tipo_documento == 'FV',
                        SiigoMovimiento.codigo_contable == '13050501',
                    ).first()
                    if existe is None:
                        return jsonify(error='Cliente no encontrado en tu cartera.'), 404
                query = query.filter(ChatCarteraHilo.identificacion == clave)
            elif vendedor_id:
                vendedor = _vendedor_chat_visible(vendedor_id=vendedor_id)
                query = query.filter(ChatCarteraHilo.vendedor_id == vendedor.id, ChatCarteraHilo.identificacion.is_(None))
            else:
                if not es_administrador():
                    vendedor = _vendedor_chat_visible()
                    query = query.filter(ChatCarteraHilo.vendedor_id == vendedor.id)
            hilos = query.order_by(ChatCarteraHilo.updated_at.desc(), ChatCarteraHilo.id.desc()).limit(100).all()
            hilos = [hilo for hilo in hilos if _hilo_chat_visible(hilo.id) is not None]
            return jsonify({'hilos': [_serializar_chat_cartera(hilo) for hilo in hilos]})

        datos = request.get_json(silent=True) or {}
        tipo = _texto(datos.get('tipo')).upper() or 'GENERAL'
        if tipo not in TIPOS_CHAT_CARTERA:
            raise ValueError('Seleccione un tipo de conversacion valido.')
        asunto = _texto(datos.get('asunto'))
        if not asunto or len(asunto) > 200:
            raise ValueError('Indique un asunto de maximo 200 caracteres.')
        mensaje = _texto(datos.get('mensaje'))
        if not mensaje or len(mensaje) > 5000:
            raise ValueError('Escriba un mensaje de maximo 5000 caracteres.')

        identificacion_raw = _texto(datos.get('identificacion'))
        cliente = _cliente_maestro_por_identificacion(identificacion_raw) if identificacion_raw else None
        vendedor = _vendedor_chat_visible(datos.get('vendedor_id'), cliente)
        clave = _nit_cartera(identificacion_raw) if identificacion_raw else None
        cliente_nombre = _texto(datos.get('cliente_nombre')) or (cliente.razon_social if cliente else None)
        hilo = ChatCarteraHilo(
            tipo=tipo,
            estado='ABIERTO',
            asunto=asunto,
            cliente_id=cliente.id if cliente else None,
            identificacion=clave,
            cliente_nombre=cliente_nombre,
            vendedor_id=vendedor.id,
            creado_por_id=current_user.id,
        )
        hilo.mensajes.append(ChatCarteraMensaje(
            remitente_id=current_user.id,
            remitente_nombre=current_user.nombre_completo,
            mensaje=mensaje,
        ))
        db.session.add(hilo)
        db.session.commit()
        return jsonify({'hilo': _serializar_chat_cartera(hilo)}), 201
    except (ValueError, PermissionError) as exc:
        db.session.rollback()
        return jsonify({'error': str(exc)}), 400 if isinstance(exc, ValueError) else 403


@contable_bp.route('/chat-cartera/<int:hilo_id>/mensajes', methods=['POST'])
@login_required
def responder_chat_cartera(hilo_id):
    try:
        _requiere_ventas()
        hilo = _hilo_chat_visible(hilo_id)
        if hilo is None:
            return jsonify(error='Conversacion no encontrada.'), 404
        datos = request.get_json(silent=True) or {}
        mensaje = _texto(datos.get('mensaje'))
        if not mensaje or len(mensaje) > 5000:
            raise ValueError('Escriba un mensaje de maximo 5000 caracteres.')
        decision = _texto(datos.get('decision')).upper()
        if decision:
            if not es_administrador():
                raise PermissionError('Solo un administrador puede registrar una decision.')
            if decision not in {'AUTORIZADO', 'RECHAZADO', 'CERRADO'}:
                raise ValueError('Seleccione una decision valida.')
            hilo.estado = decision
            if decision in {'AUTORIZADO', 'RECHAZADO'}:
                hilo.autorizado_por_id = current_user.id
                hilo.autorizado_at = datetime.utcnow()
            if decision == 'CERRADO':
                hilo.cerrado_at = datetime.utcnow()
        elif hilo.estado != 'ABIERTO':
            raise ValueError('Esta conversacion ya tiene decision registrada.')
        hilo.updated_at = datetime.utcnow()
        hilo.mensajes.append(ChatCarteraMensaje(
            remitente_id=current_user.id,
            remitente_nombre=current_user.nombre_completo,
            mensaje=mensaje,
            decision=decision or None,
        ))
        db.session.commit()
        return jsonify({'hilo': _serializar_chat_cartera(hilo)})
    except (ValueError, PermissionError) as exc:
        db.session.rollback()
        return jsonify({'error': str(exc)}), 400 if isinstance(exc, ValueError) else 403


@contable_bp.route('/seguimiento-cartera', methods=['GET', 'POST', 'PATCH'])
@login_required
def seguimiento_cartera():
    try:
        _requiere_ventas()
        datos = request.args if request.method == 'GET' else request.get_json(silent=True)
        if datos is None or not hasattr(datos, 'get'):
            raise ValueError('Debe enviar los datos del seguimiento.')
        identificacion = _texto(datos.get('identificacion'))
        clave = _nit_cartera(identificacion)
        if not clave or len(identificacion) > 50:
            raise ValueError('El cliente debe tener una identificación válida para registrar su seguimiento.')
        cliente = db.session.query(SiigoMovimiento.nombre_tercero).join(SiigoComprobante).filter(
            _alcance_movimiento(),
            SiigoMovimiento.identificacion == identificacion,
            SiigoComprobante.tipo_documento == 'FV',
            SiigoMovimiento.codigo_contable == '13050501',
        ).order_by(SiigoComprobante.fecha_elaboracion.desc(), SiigoMovimiento.id.desc()).first()
        if cliente is None:
            return jsonify({'error': 'No se encontró el cliente en la cartera SIIGO.'}), 404
        if request.method == 'PATCH':
            if type(datos.get('id')) is not int:
                raise ValueError('Indique el seguimiento del compromiso.')
            registro = SiigoSeguimientoCartera.query.filter_by(id=datos['id'], identificacion=clave).first()
            if registro is None:
                return jsonify({'error': 'Seguimiento no encontrado.'}), 404
            if not registro.fecha_compromiso:
                raise ValueError('Este seguimiento no tiene compromiso de pago.')
            if not registro.compromiso_cumplido_at:
                registro.compromiso_cumplido_at = datetime.utcnow()
                registro.compromiso_cumplido_por = current_user.nombre_completo
                db.session.commit()
            return jsonify({'seguimiento': _serializar_seguimiento_cartera(registro)})
        if request.method == 'GET':
            registros = SiigoSeguimientoCartera.query.filter_by(identificacion=clave).order_by(
                SiigoSeguimientoCartera.fecha_gestion.desc(),
                SiigoSeguimientoCartera.created_at.desc(), SiigoSeguimientoCartera.id.desc(),
            ).all()
            return jsonify({'seguimientos': [_serializar_seguimiento_cartera(item) for item in registros]})

        hoy = datetime.now(ZoneInfo('America/Bogota')).date()
        fecha_gestion = _fecha(datos.get('fecha_gestion') or hoy)
        if fecha_gestion > hoy:
            raise ValueError('La fecha de gestión no puede ser futura.')
        medio_raw = datos.get('medio')
        medios = medio_raw if isinstance(medio_raw, list) else re.split(r'[,;]+', _texto(medio_raw))
        medios = [_texto(item).upper() for item in medios if _texto(item)]
        medios_validos = {'LLAMADA', 'WHATSAPP', 'CORREO', 'VISITA', 'OTRO'}
        if not medios or any(medio not in medios_validos for medio in medios):
            raise ValueError('Seleccione el medio de contacto.')
        medio = ', '.join(dict.fromkeys(medios))
        observaciones = _texto(datos.get('observaciones'))
        if not observaciones or len(observaciones) > 5000:
            raise ValueError('Describa la gestión realizada, con un máximo de 5000 caracteres.')
        contacto = _texto(datos.get('contacto'))
        if len(contacto) > 200:
            raise ValueError('El contacto no puede superar 200 caracteres.')
        estado_gestion = _texto(datos.get('estado_gestion')).upper()
        if estado_gestion not in ESTADOS_GESTION_CARTERA:
            raise ValueError('Seleccione el estado del seguimiento de cartera.')
        fechas = {}
        for campo in ('fecha_compromiso', 'proximo_seguimiento'):
            fechas[campo] = _fecha(datos[campo]) if datos.get(campo) else None
            if fechas[campo] and fechas[campo] < fecha_gestion:
                raise ValueError('El compromiso y el próximo seguimiento no pueden ser anteriores a la gestión.')
        valor = None
        if datos.get('valor_compromiso') not in (None, ''):
            valor = _decimal(datos['valor_compromiso'])
            if not valor.is_finite() or valor <= 0 or valor >= Decimal('10000000000000000'):
                raise ValueError('El valor del compromiso debe ser positivo y menor a 10.000.000.000.000.000.')
            if valor != valor.quantize(Decimal('0.01')):
                raise ValueError('El valor del compromiso admite hasta dos decimales.')
            if not fechas['fecha_compromiso']:
                raise ValueError('Indique la fecha del compromiso de pago.')
        if estado_gestion == 'CON_COMPROMISO' and (not fechas['fecha_compromiso'] or valor is None):
            raise ValueError('Para el estado Con compromiso indique fecha y valor del compromiso.')
        identificaciones = [clave]
        if datos.get('alcance') == 'grupo':
            solicitadas = {_nit_cartera(item) for item in (datos.get('identificaciones_grupo') or [])}
            solicitadas.discard('')
            clientes_maestro = _clientes_informe()
            grupos = _clientes_agrupados_por_responsable(clientes_maestro)
            grupo = grupos.get(clave)
            permitidas = {empresa['identificacion'] for empresa in grupo['empresas']} if grupo else {clave}
            identificaciones = sorted(solicitadas & permitidas)
            if clave not in identificaciones:
                identificaciones.insert(0, clave)
            if len(identificaciones) < 2:
                raise ValueError('El responsable no tiene otras empresas visibles para registrar un compromiso agrupado.')
        nombres_cliente = {
            _nit_cartera(row.identificacion): row.nombre_tercero
            for row in db.session.query(SiigoMovimiento.identificacion, func.max(SiigoMovimiento.nombre_tercero).label('nombre_tercero'))
            .join(SiigoComprobante)
            .filter(_alcance_movimiento(), SiigoMovimiento.identificacion.in_(identificaciones),
                    SiigoComprobante.tipo_documento == 'FV', SiigoMovimiento.codigo_contable == '13050501')
            .group_by(SiigoMovimiento.identificacion).all()
        }
        registros = []
        for identificacion_registro in identificaciones:
            registro = SiigoSeguimientoCartera(
                identificacion=identificacion_registro,
                cliente_nombre=nombres_cliente.get(identificacion_registro) or ('Sin nombre' if identificacion_registro == clave else identificacion_registro),
                fecha_gestion=fecha_gestion, medio=medio, estado_gestion=estado_gestion, contacto=contacto or None,
                observaciones=observaciones, valor_compromiso=valor,
                usuario_id=current_user.id, usuario_nombre=current_user.nombre_completo,
                **fechas,
            )
            registros.append(registro)
            db.session.add(registro)
        db.session.commit()
        return jsonify({
            'seguimiento': _serializar_seguimiento_cartera(registros[0]),
            'seguimientos': [_serializar_seguimiento_cartera(item) for item in registros],
        }), 201
    except (ValueError, PermissionError) as exc:
        db.session.rollback()
        return jsonify({'error': str(exc)}), 400 if isinstance(exc, ValueError) else 403


def _seguimiento_cartera_visible(seguimiento_id):
    _requiere_ventas()
    registro = db.session.get(SiigoSeguimientoCartera, seguimiento_id)
    if registro is None:
        return None
    if not es_administrador() and registro.identificacion not in {
        _nit_cartera(c.nit) for c in clientes_visibles() if c.nit
    }:
        return None
    return registro


@contable_bp.route('/seguimiento-cartera/<int:seguimiento_id>/comprobantes', methods=['POST'])
@login_required
def subir_comprobantes_cartera(seguimiento_id):
    try:
        registro = _seguimiento_cartera_visible(seguimiento_id)
        if registro is None:
            return jsonify(error='Seguimiento no encontrado.'), 404
        archivos = request.files.getlist('archivos')
        if not 1 <= len(archivos) <= 5:
            raise ValueError('Seleccione entre 1 y 5 comprobantes de pago.')
        for archivo in archivos:
            nombre = secure_filename(archivo.filename or '')
            contenido = archivo.read(10 * 1024 * 1024 + 1)
            if not nombre or len(nombre) > 255 or not contenido:
                raise ValueError('El comprobante debe tener nombre y contenido.')
            if len(contenido) > 10 * 1024 * 1024:
                raise ValueError('Cada comprobante puede pesar máximo 10 MB.')
            extension = nombre.rsplit('.', 1)[-1].lower()
            mime = None
            if extension == 'pdf' and contenido.startswith(b'%PDF-'):
                mime = 'application/pdf'
            elif extension in {'jpg', 'jpeg'} and contenido.startswith(b'\xff\xd8\xff'):
                mime = 'image/jpeg'
            elif extension == 'png' and contenido.startswith(b'\x89PNG\r\n\x1a\n'):
                mime = 'image/png'
            elif extension == 'webp' and contenido[:4] == b'RIFF' and contenido[8:12] == b'WEBP':
                mime = 'image/webp'
            if mime is None:
                raise ValueError('Use un comprobante PDF o una imagen JPG, PNG o WebP válida.')
            db.session.add(SiigoComprobantePago(seguimiento=registro, nombre=nombre,
                mime_type=mime, tamano_bytes=len(contenido), contenido=contenido,
                usuario_id=current_user.id, usuario_nombre=current_user.nombre_completo))
        db.session.commit()
        return jsonify(seguimiento=_serializar_seguimiento_cartera(registro)), 201
    except (ValueError, PermissionError) as exc:
        db.session.rollback()
        return jsonify(error=str(exc)), 403 if isinstance(exc, PermissionError) else 400


@contable_bp.route('/seguimiento-cartera/comprobantes/<int:adjunto_id>', methods=['GET'])
@login_required
def ver_comprobante_cartera(adjunto_id):
    adjunto = db.session.get(SiigoComprobantePago, adjunto_id)
    if adjunto is None or _seguimiento_cartera_visible(adjunto.seguimiento_id) is None:
        return jsonify(error='Comprobante no encontrado.'), 404
    response = send_file(BytesIO(adjunto.contenido), mimetype=adjunto.mime_type,
                         download_name=adjunto.nombre, as_attachment=request.args.get('descargar') == '1',
                         max_age=0)
    response.headers['Cache-Control'] = 'private, no-store'
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['Content-Security-Policy'] = "sandbox; default-src 'none'"
    return response


def _clientes_facturas_vencidas(cartera_clientes, vendedores, estado_facturas='todos'):
    if estado_facturas not in {'todos', 'vencidos', 'por_vencer'}:
        raise ValueError('Seleccione Todos, Solo vencidos o Por vencer.')
    resultado = []
    for cliente in cartera_clientes:
        facturas = sorted(
            (factura for factura in cliente['facturas']
             if factura['saldo'] > 0
             and (estado_facturas == 'todos'
                  or (estado_facturas == 'vencidos' and factura['dias_vencido'] > 0)
                  or (estado_facturas == 'por_vencer' and factura['dias_vencido'] <= 0))),
            key=lambda factura: (-factura['dias_vencido'], factura['referencia']),
        )
        if not facturas:
            continue
        resultado.append({
            'vendedor': vendedores.get(_nit_cartera(cliente['identificacion']), 'Sin vendedor asignado'),
            'identificacion': cliente['identificacion'],
            'cliente': cliente['cliente'],
            'cantidad_facturas': len(facturas),
            'total_cliente': sum((factura['saldo'] for factura in facturas), Decimal('0')),
            'total_vencido': sum((f['saldo'] for f in facturas if f['dias_vencido'] > 0), Decimal('0')),
            'facturas': facturas,
        })
    resultado.sort(key=lambda cliente: (
        -cliente['cantidad_facturas'],
        -cliente['facturas'][0]['dias_vencido'] if cliente['cantidad_facturas'] == 1 else 0,
        -cliente['total_vencido'],
        _normalizar(cliente['cliente']), cliente['identificacion'],
    ))
    return resultado


def _excel_facturas_vencidas(clientes, fecha_corte, movimientos_sin_asignar=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    cantidad = max((cliente['cantidad_facturas'] for cliente in clientes), default=0)
    if 4 + cantidad * 3 > 16384:
        raise ValueError('El detalle supera el límite de columnas de Excel; filtre por cliente.')
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Facturas vencidas'
    hoja.append([f'Facturas vencidas y por vencer al {fecha_corte.isoformat()}'])
    hoja.append(['Valor = débitos menos créditos de cartera vinculados a cada factura, de cualquier tipo de comprobante, a la fecha de corte. '
                 'Facturas con saldo mayor que cero según el filtro. Días negativos: por vencer; cero: vence hoy. '
                 'Orden: cantidad, luego total vencido; clientes con una factura: de mayor a menor días de vencimiento. '
                 'Sin vencimiento SIIGO se usa la fecha de factura. Vendedor actual de la ficha comercial.'])
    encabezados = ['Vendedor', 'Cliente', 'Cantidad', 'Valor total cliente']
    for indice in range(1, cantidad + 1):
        encabezados.extend([f'N.º factura {indice}', f'Días vencida {indice}', f'Valor {indice}'])
    hoja.append(encabezados)
    for cliente in clientes:
        fila = [cliente['vendedor'], cliente['cliente'], cliente['cantidad_facturas'], cliente['total_cliente']]
        for factura in cliente['facturas']:
            fila.extend([factura['referencia'], factura['dias_vencido'], factura['saldo']])
        hoja.append(fila)
        for celda in hoja[hoja.max_row]:
            if isinstance(celda.value, str):
                celda.data_type = 's'
            if celda.column >= 4 and celda.column % 3 == 1:
                celda.number_format = '#,##0.00'
    for celda in hoja[3]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='245B45')
        hoja.column_dimensions[get_column_letter(celda.column)].width = 20
    hoja.column_dimensions['A'].width = 30
    hoja.column_dimensions['B'].width = 45
    hoja.freeze_panes = 'E4'
    hoja.auto_filter.ref = f'A3:{get_column_letter(len(encabezados))}{hoja.max_row}'
    cruces = libro.create_sheet('Cruces por factura')
    cruces.append(['Cliente', 'Factura', 'Comprobante', 'Secuencia', 'Fecha', 'Débito', 'Crédito'])
    for cliente in clientes:
        for factura in cliente['facturas']:
            for movimiento in factura.get('movimientos', []):
                cruces.append([cliente['cliente'], factura['referencia'], movimiento['comprobante'],
                               movimiento['secuencia'], movimiento['fecha'], movimiento['debito'], movimiento['credito']])
                for celda in cruces[cruces.max_row]:
                    if isinstance(celda.value, str):
                        celda.data_type = 's'
                for columna in (6, 7):
                    cruces.cell(cruces.max_row, columna).number_format = '#,##0.00'
    cruces.freeze_panes = 'C2'
    cruces.auto_filter.ref = cruces.dimensions
    for columna in 'ABCDEFG':
        cruces.column_dimensions[columna].width = 35 if columna == 'A' else 22
    if movimientos_sin_asignar:
        pendientes = libro.create_sheet('Por conciliar')
        pendientes.append(['Comprobante', 'Fecha', 'Identificación', 'Factura', 'Valor sin aplicar', 'Motivo'])
        for movimiento in movimientos_sin_asignar:
            pendientes.append([movimiento[clave] for clave in (
                'comprobante', 'fecha', 'identificacion', 'referencia', 'valor', 'motivo',
            )])
            for celda in pendientes[pendientes.max_row]:
                if isinstance(celda.value, str):
                    celda.data_type = 's'
            pendientes.cell(pendientes.max_row, 5).number_format = '#,##0.00'
        pendientes.freeze_panes = 'A2'
        pendientes.auto_filter.ref = pendientes.dimensions
        for columna in 'ABCDEF':
            pendientes.column_dimensions[columna].width = 28 if columna != 'F' else 50
    archivo = BytesIO()
    libro.save(archivo)
    libro.close()
    archivo.seek(0)
    return send_file(archivo, as_attachment=True,
                     download_name=f'facturas_vencidas_{fecha_corte.isoformat()}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@contable_bp.route('/cartera-dinamica', methods=['GET'])
@login_required
def cartera_dinamica():
    try:
        _requiere_ventas()
        fecha_corte = _fecha(request.args.get('fecha_corte') or date.today().isoformat())
        desde = _fecha(request.args.get('desde')) if request.args.get('desde') else None
        hasta = _fecha(request.args.get('hasta')) if request.args.get('hasta') else None
        cliente = _texto(request.args.get('cliente'))
        if _normalizar(cliente) in {'todos', 'todos los clientes'}:
            cliente = ''
        if desde and hasta and desde > hasta:
            raise ValueError('La fecha inicial no puede ser posterior a la fecha final.')
        facturas = {}
        consulta_facturas = db.session.query(SiigoComprobante, SiigoMovimiento).join(SiigoMovimiento).filter(
            _alcance_movimiento(),
            SiigoComprobante.tipo_documento == 'FV',
            SiigoMovimiento.codigo_contable == '13050501',
            SiigoComprobante.fecha_elaboracion <= fecha_corte,
            SiigoMovimiento.debito != SiigoMovimiento.credito,
        )
        if desde:
            consulta_facturas = consulta_facturas.filter(SiigoComprobante.fecha_elaboracion >= desde)
        if hasta:
            consulta_facturas = consulta_facturas.filter(SiigoComprobante.fecha_elaboracion <= hasta)
        if cliente:
            like = f'%{cliente}%'
            consulta_facturas = consulta_facturas.filter(or_(
                SiigoMovimiento.identificacion.ilike(like),
                SiigoMovimiento.nombre_tercero.ilike(like),
            ))
        lineas_factura = consulta_facturas.all()
        for comprobante, movimiento in lineas_factura:
            referencia = _referencia_movimiento(comprobante, movimiento)
            if not _es_base_factura(comprobante, referencia):
                continue
            referencia = referencia or f'FV-{comprobante.codigo_comprobante}-{comprobante.numero_comprobante}'
            item = facturas.setdefault(referencia, {
                'referencia': referencia,
                'cliente': movimiento.nombre_tercero or '',
                'identificacion': movimiento.identificacion or '',
                'fecha_factura': comprobante.fecha_elaboracion,
                'fecha_vencimiento': _fecha_vencimiento(movimiento.detalle),
                'valor_factura': Decimal('0'),
                'recaudado': Decimal('0'),
                'ajustes_ac': Decimal('0'),
                'notas_credito': Decimal('0'),
                'notas_debito': Decimal('0'),
                'otros_movimientos': Decimal('0'),
                'pagos': [],
                'cancelaciones': [],
                'movimientos': [],
            })
            item['valor_factura'] += movimiento.debito - movimiento.credito
            item['movimientos'].append(_movimiento_cartera(comprobante, movimiento))
            item['fecha_vencimiento'] = item['fecha_vencimiento'] or _fecha_vencimiento(movimiento.detalle)

        pagos_sin_factura = 0
        ajustes_ac_sin_factura = 0
        valor_ac_sin_factura = Decimal('0')
        notas_credito_sin_asignar = Decimal('0')
        notas_debito_sin_asignar = Decimal('0')
        otros_sin_asignar = Decimal('0')
        movimientos_sin_asignar = []
        for comprobante, movimiento, referencia, valor in _cancelaciones_cartera(fecha_corte):
            tipo = _tipo_cartera(comprobante.tipo_documento, comprobante.codigo_comprobante)
            item = facturas.get(referencia)
            if not _mismo_tercero_cartera(item, movimiento):
                if tipo == 'AC':
                    ajustes_ac_sin_factura += 1
                    valor_ac_sin_factura += valor
                elif tipo == 'NC':
                    notas_credito_sin_asignar += valor
                elif tipo == 'ND':
                    notas_debito_sin_asignar -= valor
                elif tipo == 'RC':
                    pagos_sin_factura += 1
                else:
                    otros_sin_asignar += valor
                movimientos_sin_asignar.append({
                    'comprobante': f'{comprobante.tipo_documento}-{comprobante.codigo_comprobante}-{comprobante.numero_comprobante}',
                    'fecha': comprobante.fecha_elaboracion.isoformat(),
                    'referencia': referencia,
                    'identificacion': movimiento.identificacion,
                    'valor': float(valor),
                    'motivo': 'Sin referencia única a factura' if not referencia else (
                        'Factura no incluida en la consulta' if item is None else 'El tercero no coincide'),
                })
                continue
            item[APLICACIONES_CARTERA.get(tipo, 'otros_movimientos')] += -valor if tipo == 'ND' else valor
            item['movimientos'].append(_movimiento_cartera(comprobante, movimiento))
            cancelacion = {'fecha': comprobante.fecha_elaboracion, 'valor': valor}
            item['cancelaciones'].append(cancelacion)
            if tipo == 'RC':
                item['pagos'].append(cancelacion)

        periodos = {}
        cartera_por_cliente = {}
        pagos_completos = []
        for item in facturas.values():
            saldo = max(item['valor_factura'] - sum(
                (movimiento['valor'] for movimiento in item['cancelaciones']), Decimal('0'),
            ), Decimal('0'))
            vencimiento = item['fecha_vencimiento'] or item['fecha_factura']
            dias_vencido = (fecha_corte - vencimiento).days
            periodo = item['fecha_factura'].strftime('%Y-%m')
            resumen = periodos.setdefault(periodo, {
                'periodo': periodo, 'facturado': Decimal('0'), 'recaudado': Decimal('0'), 'ajustes_ac': Decimal('0'), 'notas_credito': Decimal('0'), 'notas_debito': Decimal('0'), 'saldo': Decimal('0'),
                'otros_movimientos': Decimal('0'),
                'por_vencer': Decimal('0'), 'vencido_1_30': Decimal('0'), 'vencido_31_60': Decimal('0'),
                'vencido_61_90': Decimal('0'), 'vencido_91_mas': Decimal('0'), 'documentos': 0,
            })
            resumen['facturado'] += item['valor_factura']
            resumen['recaudado'] += item['recaudado']
            resumen['ajustes_ac'] += item['ajustes_ac']
            resumen['notas_credito'] += item['notas_credito']
            resumen['notas_debito'] += item['notas_debito']
            resumen['otros_movimientos'] += item['otros_movimientos']
            resumen['saldo'] += saldo
            resumen['documentos'] += 1
            if dias_vencido <= 0:
                resumen['por_vencer'] += saldo
            elif dias_vencido <= 30:
                resumen['vencido_1_30'] += saldo
            elif dias_vencido <= 60:
                resumen['vencido_31_60'] += saldo
            elif dias_vencido <= 90:
                resumen['vencido_61_90'] += saldo
            else:
                resumen['vencido_91_mas'] += saldo

            cliente_nombre = item['cliente'] or 'Sin nombre'
            cliente_clave = item['identificacion'] or _normalizar(cliente_nombre) or item['referencia']
            resumen_cliente = cartera_por_cliente.setdefault(cliente_clave, {
                'identificacion': item['identificacion'],
                'cliente': cliente_nombre,
                'facturado': Decimal('0'),
                'recaudado': Decimal('0'),
                'ajustes_ac': Decimal('0'),
                'notas_credito': Decimal('0'),
                'notas_debito': Decimal('0'),
                'otros_movimientos': Decimal('0'),
                'saldo': Decimal('0'),
                'por_vencer': Decimal('0'),
                'vencido_1_30': Decimal('0'),
                'vencido_31_60': Decimal('0'),
                'vencido_61_90': Decimal('0'),
                'vencido_91_mas': Decimal('0'),
                'facturas': [],
            })
            resumen_cliente['facturado'] += item['valor_factura']
            resumen_cliente['recaudado'] += item['recaudado']
            resumen_cliente['ajustes_ac'] += item['ajustes_ac']
            resumen_cliente['notas_credito'] += item['notas_credito']
            resumen_cliente['notas_debito'] += item['notas_debito']
            resumen_cliente['otros_movimientos'] += item['otros_movimientos']
            resumen_cliente['saldo'] += saldo
            if dias_vencido <= 0:
                resumen_cliente['por_vencer'] += saldo
            elif dias_vencido <= 30:
                resumen_cliente['vencido_1_30'] += saldo
            elif dias_vencido <= 60:
                resumen_cliente['vencido_31_60'] += saldo
            elif dias_vencido <= 90:
                resumen_cliente['vencido_61_90'] += saldo
            else:
                resumen_cliente['vencido_91_mas'] += saldo
            resumen_cliente['facturas'].append({
                'referencia': item['referencia'],
                'fecha_factura': item['fecha_factura'].isoformat(),
                'fecha_vencimiento': vencimiento.isoformat(),
                'facturado': item['valor_factura'],
                'recaudado': item['recaudado'],
                'ajustes_ac': item['ajustes_ac'],
                'notas_credito': item['notas_credito'],
                'notas_debito': item['notas_debito'],
                'otros_movimientos': item['otros_movimientos'],
                'saldo': saldo,
                'dias_vencido': dias_vencido,
                'movimientos': sorted(item['movimientos'], key=lambda mov: (mov['fecha'], mov['comprobante'], mov['secuencia'])),
            })

            if saldo == 0 and item['pagos']:
                fecha_pago_total = max(pago['fecha'] for pago in item['cancelaciones'])
                pagos_completos.append({
                    'referencia': item['referencia'], 'cliente': item['cliente'],
                    'identificacion': item['identificacion'],
                    'dias': (fecha_pago_total - item['fecha_factura']).days,
                    'fecha_pago': fecha_pago_total,
                })

        clientes_pagos = {}
        for pago in pagos_completos:
            cliente_pago = clientes_pagos.setdefault(pago['identificacion'] or pago['cliente'], {
                'cliente': pago['cliente'], 'pagos': [],
            })
            cliente_pago['pagos'].append(pago)
        analisis_pagos = []
        for cliente_pago in clientes_pagos.values():
            pagos = cliente_pago['pagos']
            rapido = min(pagos, key=lambda pago: pago['dias'])
            lento = max(pagos, key=lambda pago: pago['dias'])
            analisis_pagos.append({
                'cliente': cliente_pago['cliente'], 'facturas_pagadas': len(pagos),
                'promedio_dias': round(sum(pago['dias'] for pago in pagos) / len(pagos), 1),
                'mas_rapida': rapido['referencia'], 'dias_mas_rapida': rapido['dias'],
                'mas_lenta': lento['referencia'], 'dias_mas_lenta': lento['dias'],
            })
        analisis_pagos.sort(key=lambda item: item['promedio_dias'], reverse=True)

        totales = {
            'ventas': sum((item['facturado'] for item in periodos.values()), Decimal('0')),
            'recaudado': sum((item['recaudado'] for item in periodos.values()), Decimal('0')),
            'sin_recaudo': sum((item['saldo'] for item in cartera_por_cliente.values()), Decimal('0')),
            'cartera': sum((item['saldo'] for item in cartera_por_cliente.values()), Decimal('0')),
        }

        def serializar(items):
            return [{clave: (float(valor) if isinstance(valor, Decimal) else valor) for clave, valor in item.items()} for item in items]

        def serializar_cartera_clientes():
            resultado = []
            for cliente_cartera in sorted(cartera_por_cliente.values(), key=lambda item: item['saldo'], reverse=True):
                registro = {
                    clave: (float(valor) if isinstance(valor, Decimal) else valor)
                    for clave, valor in cliente_cartera.items()
                    if clave != 'facturas'
                }
                registro['facturas'] = serializar(sorted(
                    cliente_cartera['facturas'],
                    key=lambda item: (item['fecha_factura'], item['referencia']),
                ))
                resultado.append(registro)
            return resultado

        clientes_maestro_informe = _clientes_informe()
        vendedores_maestro = {identificacion(c.nit): c.vendedor.nombre if c.vendedor else 'Sin asignar' for c in clientes_maestro_informe if c.nit}
        for registro_cliente in cartera_por_cliente.values():
            registro_cliente['vendedor'] = vendedores_maestro.get(identificacion(registro_cliente['identificacion']), 'Sin asignar')

        if request.args.get('informe') == 'vencidas':
            vendedores_por_nit = {}
            for cliente_maestro in clientes_maestro_informe:
                clave = _nit_cartera(cliente_maestro.nit)
                if clave:
                    vendedores_por_nit.setdefault(clave, set()).add(cliente_maestro.vendedor.nombre if cliente_maestro.vendedor else 'Sin asignar')
            vendedores = {
                nit: next(iter(nombres)) if len(nombres) == 1 else 'Asignación por revisar'
                for nit, nombres in vendedores_por_nit.items()
            }
            estado_facturas = request.args.get('estado_facturas') or ('por_vencer' if request.args.get('solo_por_vencer') == '1' else 'todos')
            clientes = _clientes_facturas_vencidas(cartera_por_cliente.values(), vendedores, estado_facturas)
            agrupaciones = _clientes_agrupados_por_responsable(clientes_maestro_informe)
            claves = {_nit_cartera(c['identificacion']) for c in clientes}
            seguimientos = {}
            if claves:
                for registro in SiigoSeguimientoCartera.query.filter(SiigoSeguimientoCartera.identificacion.in_(claves)).all():
                    seguimientos.setdefault(registro.identificacion, []).append(_serializar_seguimiento_cartera(registro))
            for cliente in clientes:
                cliente['seguimientos'] = seguimientos.get(_nit_cartera(cliente['identificacion']), [])
                cliente['agrupacion_responsable'] = agrupaciones.get(_nit_cartera(cliente['identificacion']))
            if request.args.get('formato') == 'xlsx':
                return _excel_facturas_vencidas(clientes, fecha_corte, movimientos_sin_asignar)
            return jsonify({
                'fecha_corte': fecha_corte.isoformat(),
                'clientes': [dict(cliente, total_vencido=float(cliente['total_vencido']), total_cliente=float(cliente['total_cliente']),
                                  facturas=serializar(cliente['facturas'])) for cliente in clientes],
                'cantidad_clientes': len(clientes),
                'cantidad_facturas': sum(cliente['cantidad_facturas'] for cliente in clientes),
                'total_vencido': float(sum((cliente['total_vencido'] for cliente in clientes), Decimal('0'))),
                'notas_credito_sin_asignar': float(notas_credito_sin_asignar or 0),
                'notas_debito_sin_asignar': float(notas_debito_sin_asignar),
                'otros_sin_asignar': float(otros_sin_asignar),
                'movimientos_sin_asignar': movimientos_sin_asignar,
                'pagos_sin_factura': pagos_sin_factura,
                'ajustes_ac_sin_factura': ajustes_ac_sin_factura,
            })

        return jsonify({
            'fecha_corte': fecha_corte.isoformat(),
            'desde': desde.isoformat() if desde else None,
            'hasta': hasta.isoformat() if hasta else None,
            'cliente': cliente or None,
            'periodos': serializar(sorted(periodos.values(), key=lambda item: item['periodo'])),
            'cartera_clientes': serializar_cartera_clientes(),
            'totales': {clave: float(valor) for clave, valor in totales.items()},
            'pagos_clientes': analisis_pagos,
            'pagos_sin_factura': pagos_sin_factura,
            'ajustes_ac_sin_factura': ajustes_ac_sin_factura,
            'valor_ac_sin_factura': float(valor_ac_sin_factura),
            'notas_credito_sin_asignar': float(notas_credito_sin_asignar or 0),
            'notas_debito_sin_asignar': float(notas_debito_sin_asignar),
            'otros_sin_asignar': float(otros_sin_asignar),
            'movimientos_sin_asignar': movimientos_sin_asignar,
            'facturas': len(facturas),
        })
    except (ValueError, PermissionError) as exc:
        return jsonify({'error': str(exc)}), 400 if isinstance(exc, ValueError) else 403
