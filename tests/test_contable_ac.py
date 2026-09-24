"""Pruebas de cartera AC en una base PostgreSQL exclusiva para pruebas.

TEST_DATABASE_URL debe apuntar a una base cuyo nombre empiece por prevent_test_ac_.
Ejecutar: python -m unittest discover -s tests -v
"""

from datetime import date
from decimal import Decimal
from io import BytesIO
import os
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from flask import Flask, g
from flask_login import LoginManager
from openpyxl import Workbook, load_workbook
from sqlalchemy.engine import make_url

from app.models import (db, ClienteComercial, Vendedor, SiigoCarga, SiigoComprobante,
                        SiigoMovimiento, SiigoSeguimientoCartera, Usuario, Role)
from app.routes import contable


class CarteraACTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        url = os.environ.get('TEST_DATABASE_URL')
        if not url:
            raise unittest.SkipTest('Requiere TEST_DATABASE_URL de PostgreSQL')
        parsed = make_url(url)
        if not parsed.drivername.startswith('postgresql') or not parsed.database.startswith('prevent_test_ac_'):
            raise RuntimeError('Use una base PostgreSQL exclusiva prevent_test_ac_*')
        cls.app = Flask(__name__)
        cls.app.config.update(TESTING=True, SECRET_KEY='solo-pruebas', SQLALCHEMY_DATABASE_URI=url)
        db.init_app(cls.app)
        login = LoginManager(cls.app)
        login.user_loader(lambda user_id: db.session.get(Usuario, int(user_id)))
        login.unauthorized_handler(lambda: ({'error': 'Sesión requerida'}, 401))
        cls.app.register_blueprint(contable.contable_bp)
        with cls.app.app_context():
            db.create_all()

    def setUp(self):
        self.context = self.app.app_context()
        self.context.push()
        self.addCleanup(self.context.pop)
        self.addCleanup(db.session.remove)
        for model in (SiigoSeguimientoCartera, SiigoMovimiento, SiigoComprobante, SiigoCarga):
            db.session.query(model).delete()
        db.session.commit()
        # Estas pruebas verifican el consolidado contable. El aislamiento real
        # por vendedor se prueba con usuarios autenticados en test_clientes_maestro.
        for target in ('app.clientes_scope.es_administrador', 'app.routes.contable.es_administrador',
                       'app.routes.clientes_acceso.es_administrador'):
            scope_patch = patch(target, return_value=True)
            scope_patch.start()
            self.addCleanup(scope_patch.stop)
        self.permission = patch.object(contable, '_requiere_ventas')
        self.permission.start()
        self.addCleanup(self.permission.stop)

    def documento(self, tipo, numero, fecha, lineas, codigo=None):
        carga = SiigoCarga.query.first()
        if carga is None:
            carga = SiigoCarga(tipo_archivo='COMPROBANTES', nombre_archivo='fixture.xlsx', hash_archivo='fixture')
            db.session.add(carga)
            db.session.flush()
        doc = SiigoComprobante(tipo_documento=tipo, codigo_comprobante=codigo or ('2' if tipo == 'FV' else '1'),
                               numero_comprobante=str(numero), fecha_elaboracion=date.fromisoformat(fecha), carga_id=carga.id)
        for index, fields in enumerate(lineas, 1):
            values = dict(codigo_contable='13050501', cuenta_contable='Clientes', identificacion='9001',
                          nombre_tercero='Cliente prueba', debito=Decimal('0'), credito=Decimal('0'),
                          detalle='FV-2-1 Cuota: 1 Fecha: 31/01/2026', descripcion='Clientes nacionales')
            values.update(fields)
            doc.movimientos.append(SiigoMovimiento(secuencia=index, **values))
        db.session.add(doc)
        db.session.commit()

    def factura_y_recibo(self):
        self.documento('FV', 1, '2026-01-01', [{'debito': Decimal('1000')}])
        self.documento('RC', 1, '2026-02-01', [{'credito': Decimal('900'), 'descripcion': 'Abono FV-2-1'}])

    def cartera(self, corte='2026-02-28'):
        with self.app.test_request_context(query_string={'fecha_corte': corte}):
            return contable.cartera_dinamica.__wrapped__().get_json()

    def test_retenciones_ac_cancelan_saldo_sin_inflar_recaudado(self):
        self.factura_y_recibo()
        self.documento('AC', 1, '2026-02-10', [
            {'credito': Decimal('70')}, {'credito': Decimal('30')},
            {'codigo_contable': '13551501', 'debito': Decimal('100')},
        ])
        data = self.cartera()
        for result in (data['periodos'][0], data['cartera_clientes'][0], data['cartera_clientes'][0]['facturas'][0]):
            self.assertEqual(result['facturado'], 1000)
            self.assertEqual(result['recaudado'], 900)
            self.assertEqual(result['ajustes_ac'], 100)
            self.assertEqual(result['saldo'], 0)
        self.assertEqual(data['cartera_clientes'][0]['vencido_1_30'], 0)
        self.assertEqual(data['pagos_clientes'][0]['promedio_dias'], 40)

    def test_ac_respeta_corte_y_reversos_reabren_saldo(self):
        self.factura_y_recibo()
        self.documento('AC', 1, '2026-02-10', [{'credito': Decimal('100')}])
        self.documento('AC', 2, '2026-02-20', [{'debito': Decimal('30')}])
        for corte, ajuste, saldo in [('2026-02-09', 0, 100), ('2026-02-10', 100, 0), ('2026-02-28', 70, 30)]:
            with self.subTest(corte=corte):
                data = self.cartera(corte)
                cliente = data['cartera_clientes'][0]
                self.assertEqual(cliente['ajustes_ac'], ajuste)
                self.assertEqual(cliente['saldo'], saldo)
                self.assertEqual(cliente['vencido_1_30'], saldo)

    def test_ac_sin_referencia_o_de_otro_tercero_no_cancela_factura(self):
        self.factura_y_recibo()
        self.documento('AC', 1, '2026-02-10', [
            {'credito': Decimal('10'), 'detalle': 'Sin referencia'},
            {'credito': Decimal('20'), 'detalle': 'FV-2-999 Cuota: 1'},
            {'credito': Decimal('30'), 'identificacion': 'OTRO'},
        ])
        data = self.cartera()
        self.assertEqual(data['cartera_clientes'][0]['saldo'], 100)
        self.assertEqual(data['ajustes_ac_sin_factura'], 3)
        self.assertEqual(data['valor_ac_sin_factura'], 60)

    def test_recibo_sin_referencia_se_aplica_a_facturas_abiertas_del_mismo_cliente(self):
        self.documento('FV', 1, '2026-01-01', [{'debito': Decimal('1000')}])
        self.documento('FV', 2, '2026-01-15', [{'debito': Decimal('500'), 'detalle': 'FV-2-2 Cuota: 1 Fecha: 15/02/2026'}])
        self.documento('RC', 1, '2026-02-01', [{'credito': Decimal('1200'), 'detalle': 'Pago cliente', 'descripcion': 'Pago cliente'}])
        data = self.cartera()
        facturas = {f['referencia']: f for c in data['cartera_clientes'] for f in c['facturas']}
        self.assertEqual(facturas['FV-2-1']['saldo'], 0)
        self.assertEqual(facturas['FV-2-2']['saldo'], 300)
        self.assertEqual(data['cartera_clientes'][0]['saldo'], 300)
        self.assertEqual(data['pagos_sin_factura'], 0)

    def test_ac_puede_usar_descripcion_y_prioriza_detalle(self):
        self.factura_y_recibo()
        self.documento('AC', 1, '2026-02-10', [
            {'credito': Decimal('30'), 'detalle': '', 'descripcion': 'Ajuste FV-2-1'},
            {'credito': Decimal('20'), 'descripcion': 'Texto FV-2-999'},
        ])
        self.assertEqual(self.cartera()['cartera_clientes'][0]['saldo'], 50)

    def test_ac_por_si_solo_no_es_pago_en_efectivo(self):
        self.documento('FV', 1, '2026-01-01', [{'debito': Decimal('1000')}])
        self.documento('AC', 1, '2026-02-10', [{'credito': Decimal('1000')}])
        data = self.cartera()
        self.assertEqual(data['cartera_clientes'][0]['saldo'], 0)
        self.assertEqual(data['cartera_clientes'][0]['recaudado'], 0)
        self.assertEqual(data['pagos_clientes'], [])

    def test_cc_ac_cancela_retenciones_consultores_y_respeta_corte(self):
        tercero = {'identificacion': '900416292', 'nombre_tercero': 'CONSULTORES EN GESTION HUMANA SAS',
                   'detalle': 'FV-2-3658 Cuota: 1 Fecha: 26/02/2025'}
        self.documento('FV', 3658, '2025-02-11', [dict(tercero, debito=Decimal('5408000'))])
        self.documento('RC', 1, '2025-03-19', [dict(tercero, credito=Decimal('5247559'), descripcion='FV-2-3658')])
        self.documento('CC', 827, '2025-03-19', [
            dict(tercero, credito=Decimal('108160')),
            dict(tercero, codigo_contable='13551515', debito=Decimal('108160')),
            dict(tercero, credito=Decimal('52281')),
            dict(tercero, codigo_contable='13551805', debito=Decimal('52281')),
        ], codigo='AC')
        # Un CC de otro código también cruza, pero se desglosa separado de los AC.
        self.documento('CC', 1, '2025-03-19', [dict(tercero, credito=Decimal('999'))])
        antes = self.cartera('2025-03-18')['cartera_clientes'][0]
        self.assertEqual((antes['saldo'], antes['ajustes_ac']), (5408000, 0))
        despues = self.cartera('2025-03-19')['cartera_clientes'][0]
        self.assertEqual((despues['recaudado'], despues['ajustes_ac'], despues['saldo']), (5247559, 160441, 0))
        self.assertEqual(despues['otros_movimientos'], 999)
        with self.app.test_request_context(query_string={'informe': 'vencidas', 'fecha_corte': '2025-03-19'}):
            informe = contable.cartera_dinamica.__wrapped__().get_json()
        self.assertEqual(informe['clientes'], [])
        with self.app.test_request_context(query_string={'tipo': 'AC', 'cliente': '900416292'}):
            ajustes = contable.consultar_comprobantes.__wrapped__().get_json()['comprobantes']
        self.assertEqual([(item['tipo'], item['codigo'], item['numero']) for item in ajustes], [('CC', 'AC', '827')])

    def test_reimportar_recupera_cc_ac_omitidos_sin_duplicar(self):
        libro = load_workbook(BytesIO(self.excel()))
        for fila in libro.active:
            if fila[0].value == 'Comprobante: AC-1-1':
                fila[0].value = 'Comprobante: CC-AC-1'
        archivo = BytesIO()
        libro.save(archivo)
        libro.close()
        contenido = archivo.getvalue()
        with patch.object(contable, '_comprobante_relevante', side_effect=lambda doc, cols: doc['tipo'] in {'FV', 'RC', 'NC', 'ND', 'AC'}):
            data, status = self.cargar(contenido)
        self.assertEqual((status, data['comprobantes']), (200, 2))
        self.assertEqual(self.cartera()['cartera_clientes'][0]['saldo'], 100)
        data, status = self.cargar(contenido)
        self.assertEqual((status, data['comprobantes']), (200, 1))
        self.assertEqual(SiigoComprobante.query.filter_by(tipo_documento='CC', codigo_comprobante='AC').count(), 1)
        self.assertEqual(self.cartera()['cartera_clientes'][0]['saldo'], 0)
        self.assertEqual(SiigoCarga.query.count(), 1)
        self.assertEqual(SiigoCarga.query.one().registros_omitidos, 0)
        data, status = self.cargar(contenido)
        self.assertEqual(status, 400)
        self.assertEqual(SiigoComprobante.query.count(), 3)

    def test_comparativo_coincide_con_cartera_y_corte(self):
        self.factura_y_recibo()
        self.documento('AC', 1, '2026-02-10', [{'credito': Decimal('100')}])
        for corte, saldo in [('2026-02-09', 100), ('2026-02-28', 0)]:
            with self.app.test_request_context(query_string={
                'periodo_a_desde': '2025-01-01', 'periodo_a_hasta': '2025-12-31',
                'periodo_b_desde': '2026-01-01', 'periodo_b_hasta': '2026-01-31',
                'fecha_corte_cartera': corte,
            }):
                data = contable.comparativo_clientes.__wrapped__().get_json()
            self.assertEqual(data['totales_nuevos']['cartera'], saldo)
            self.assertEqual(data['totales_nuevos']['facturacion'], 1000)
            self.assertEqual(data['nuevos'][0]['primera_factura_fecha'], '2026-01-01')
            self.assertEqual(data['nuevos'][0]['primera_factura_valor'], 1000)
            self.assertEqual(data['nuevos'][0]['ultima_factura_fecha'], '2026-01-01')
            self.assertEqual(data['nuevos'][0]['ultima_factura_valor'], 1000)

    def test_comparativo_normaliza_identificacion_decimal_texto(self):
        self.documento('FV', 1, '2024-09-01', [{
            'identificacion': '900532173.0',
            'nombre_tercero': 'SORING CLINICA',
            'debito': Decimal('100'),
        }])
        self.documento('FV', 2, '2026-03-01', [{
            'identificacion': '900532173',
            'nombre_tercero': 'SORING CLINICA',
            'debito': Decimal('200'),
        }])
        with self.app.test_request_context(query_string={
            'periodo_a_desde': '2024-01-01', 'periodo_a_hasta': '2025-12-31',
            'periodo_b_desde': '2026-01-01', 'periodo_b_hasta': '2026-09-15',
        }):
            comparativo = contable.comparativo_clientes.__wrapped__().get_json()
        self.assertEqual(comparativo['nuevos'], [])
        self.assertEqual(comparativo['no_volvieron'], [])

    def test_nc_rc_ac_cancelan_factura_y_reversos_reabren_saldo_al_corte(self):
        self.factura_y_recibo()
        self.documento('AC', 1, '2026-02-05', [{'credito': Decimal('30')}])
        self.documento('NC', 1, '2026-02-06', [
            {'credito': Decimal('20')},
            {'codigo_contable': '41350501', 'debito': Decimal('20')},
        ])
        self.documento('NC', 2, '2026-02-10', [{'credito': Decimal('50')}])
        self.documento('NC', 3, '2026-02-20', [{'debito': Decimal('10')}])
        self.documento('ND', 1, '2026-02-21', [{'debito': Decimal('5')}])
        self.documento('RC', 2, '2026-02-22', [{'debito': Decimal('5'), 'descripcion': 'Reverso FV-2-1'}])
        for corte, saldo in [('2026-02-09', 50), ('2026-02-10', 0), ('2026-02-20', 10),
                             ('2026-02-21', 15), ('2026-02-22', 20)]:
            with self.subTest(corte=corte):
                cartera = self.cartera(corte)
                self.assertEqual(cartera['cartera_clientes'][0]['saldo'], saldo)
                self.assertEqual(cartera['periodos'][0]['saldo'], saldo)
                self.assertEqual(cartera['notas_credito_sin_asignar'], 0)
                with self.app.test_request_context(query_string={'informe': 'vencidas', 'fecha_corte': corte}):
                    informe = contable.cartera_dinamica.__wrapped__().get_json()
                self.assertEqual(informe['cantidad_facturas'], int(saldo > 0))
                self.assertEqual(informe['total_vencido'], saldo)
                with self.app.test_request_context(query_string={
                    'periodo_a_desde': '2025-01-01', 'periodo_a_hasta': '2025-12-31',
                    'periodo_b_desde': '2026-01-01', 'periodo_b_hasta': '2026-01-31',
                    'fecha_corte_cartera': corte,
                }):
                    comparativo = contable.comparativo_clientes.__wrapped__().get_json()
                self.assertEqual(comparativo['totales_nuevos']['cartera'], saldo)
        detalle = self.cartera()['cartera_clientes'][0]['facturas'][0]
        self.assertEqual((detalle['recaudado'], detalle['ajustes_ac'], detalle['notas_credito'], detalle['notas_debito']),
                         (895, 30, 60, 5))

    def test_movimientos_sin_factura_o_ambiguos_no_se_descuentan(self):
        self.factura_y_recibo()
        self.documento('NC', 1, '2026-02-05', [
            {'credito': Decimal('10'), 'detalle': 'NC-1-1'},
            {'credito': Decimal('20'), 'identificacion': 'OTRO'},
            {'credito': Decimal('30'), 'detalle': 'FV-2-999'},
            {'credito': Decimal('40'), 'detalle': 'FV-2-1 y FV-2-2'},
        ])
        data = self.cartera()
        self.assertEqual(data['cartera_clientes'][0]['saldo'], 100)
        self.assertEqual(data['notas_credito_sin_asignar'], 100)
        self.assertEqual(len(data['movimientos_sin_asignar']), 4)
        with self.app.test_request_context(query_string={
            'informe': 'vencidas', 'fecha_corte': '2026-02-28', 'formato': 'xlsx',
        }):
            response = contable.cartera_dinamica.__wrapped__()
            response.direct_passthrough = False
            libro = load_workbook(BytesIO(response.get_data()))
            self.assertEqual(libro['Por conciliar'].max_row, 5)
            self.assertEqual(libro['Facturas vencidas']['G4'].value, 100)
            libro.close()
            response.close()

    def test_referencia_en_detalle_rc_nit_formateado_y_puntuacion_ac(self):
        self.documento('FV', 1, '2026-01-01', [{'debito': Decimal('100')}])
        self.documento('RC', 1, '2026-02-01', [{
            'debito': Decimal('10'), 'credito': Decimal('80'),
            'descripcion': 'Clientes nacionales', 'detalle': 'FV-2-1 Cuota: 1',
        }])
        self.documento('AC', 1, '2026-02-02', [{
            'credito': Decimal('30'), 'identificacion': '9.001-2', 'detalle': 'Ajuste (FV-2-1).',
        }])
        data = self.cartera()
        self.assertEqual(data['cartera_clientes'][0]['saldo'], 0)
        self.assertEqual(data['cartera_clientes'][0]['recaudado'], 70)
        self.assertEqual(data['ajustes_ac_sin_factura'], 0)
        with self.app.test_request_context(query_string={
            'periodo_a_desde': '2025-01-01', 'periodo_a_hasta': '2025-12-31',
            'periodo_b_desde': '2026-01-01', 'periodo_b_hasta': '2026-01-31',
            'fecha_corte_cartera': '2026-02-28',
        }):
            comparativo = contable.comparativo_clientes.__wrapped__().get_json()
        self.assertEqual(comparativo['totales_nuevos']['cartera'], 0)

    def test_informe_vencidas_cuenta_facturas_y_respeta_saldo_corte_y_vendedor(self):
        vendedor = Vendedor(nombre='Vendedora de prueba')
        cliente = ClienteComercial(nit='9.001-2', razon_social='Cliente prueba', vendedor=vendedor)
        db.session.add(cliente)
        db.session.commit()
        try:
            for numero, nit, valor, vencimiento in [
                (1, '9001', '1000', '31/01/2026'),
                (2, '9001', '50', '15/01/2026'),
                (3, '9001', '500', '28/02/2026'),  # Vence hoy: no vencida.
                (4, '9001', '500', '01/03/2026'),  # Por vencer.
                (5, '9001', '500', '31/01/2026'),  # Pagada.
                (6, '9002', '9000', '31/01/2026'),  # Mayor valor, menor cantidad.
            ]:
                self.documento('FV', numero, '2026-01-01', [{
                    'identificacion': nit, 'debito': Decimal(valor),
                    'detalle': f'FV-2-{numero} Cuota: 1 Fecha: {vencimiento}',
                }])
            # Dos líneas de una misma factura no deben aumentar la cantidad.
            self.documento('FV', 7, '2026-01-01', [
                {'debito': Decimal('100'), 'detalle': 'FV-2-7 Fecha: 31/01/2026'},
                {'debito': Decimal('50'), 'detalle': 'FV-2-7 Fecha: 31/01/2026'},
            ])
            self.documento('RC', 1, '2026-02-01', [
                {'credito': Decimal('900'), 'descripcion': 'Abono FV-2-1'},
                {'credito': Decimal('500'), 'descripcion': 'Abono FV-2-5'},
            ])
            self.documento('AC', 1, '2026-02-10', [{'credito': Decimal('30')}])
            self.documento('RC', 2, '2026-03-01', [{'credito': Decimal('70'), 'descripcion': 'Abono FV-2-1'}])
            with self.app.test_request_context(query_string={'informe': 'vencidas', 'fecha_corte': '2026-02-28'}):
                data = contable.cartera_dinamica.__wrapped__().get_json()
            self.assertEqual([c['identificacion'] for c in data['clientes']], ['9001', '9002'])
            primero = data['clientes'][0]
            self.assertEqual(primero['vendedor'], 'Vendedora de prueba')
            self.assertEqual(primero['cantidad_facturas'], 5)
            self.assertEqual(primero['total_vencido'], 270)
            self.assertEqual(primero['total_cliente'], 1270)
            self.assertEqual([f['referencia'] for f in primero['facturas']], ['FV-2-2', 'FV-2-1', 'FV-2-7', 'FV-2-3', 'FV-2-4'])
            self.assertEqual(primero['facturas'][1]['saldo'], 70)
            self.assertEqual(primero['facturas'][1]['dias_vencido'], 28)
            self.assertEqual(data['cantidad_facturas'], 6)
            self.assertEqual(data['clientes'][1]['vendedor'], 'Sin vendedor asignado')
            with self.app.test_request_context(query_string={
                'informe': 'vencidas', 'fecha_corte': '2026-02-28', 'cliente': '9002',
            }):
                filtrado = contable.cartera_dinamica.__wrapped__().get_json()
            self.assertEqual(filtrado['cantidad_clientes'], 1)
            self.assertEqual(filtrado['clientes'][0]['identificacion'], '9002')
        finally:
            db.session.delete(cliente)
            db.session.delete(vendedor)
            db.session.commit()

    def test_informe_vencidas_excel_y_cliente_sin_vencidas(self):
        self.factura_y_recibo()
        with self.app.test_request_context(query_string={
            'informe': 'vencidas', 'fecha_corte': '2026-02-28', 'formato': 'xlsx',
        }):
            response = contable.cartera_dinamica.__wrapped__()
            response.direct_passthrough = False
            libro = load_workbook(BytesIO(response.get_data()))
            self.assertEqual(list(libro.active.values)[3],
                             ('Sin vendedor asignado', 'Cliente prueba', 1, 100, 'FV-2-1', 28, 100))
            self.assertEqual(libro.active.freeze_panes, 'E4')
            libro.close()
            response.close()
        with self.app.test_request_context(query_string={'informe': 'vencidas', 'fecha_corte': '2026-01-31'}):
            data = contable.cartera_dinamica.__wrapped__().get_json()
        self.assertEqual(data['cantidad_facturas'], 1)
        self.assertEqual(data['total_vencido'], 0)

    def usuario_seguimiento(self):
        usuario = Usuario.query.filter_by(usuario='gestor_test').first()
        if not usuario:
            role = Role(nombre='Gestor prueba sin permisos')
            usuario = Usuario(usuario='gestor_test', nombre_completo='Gestora cartera',
                              email='gestor_test@example.test', password_hash='no-login', role=role)
            db.session.add(usuario)
            db.session.commit()
        return usuario

    def test_seguimiento_persiste_historial_por_cliente_y_autor_real(self):
        self.factura_y_recibo()
        self.documento('FV', 2, '2026-01-01', [{'identificacion': '9002', 'detalle': 'FV-2-2', 'debito': Decimal('10')}])
        usuario = self.usuario_seguimiento()
        http = self.app.test_client()
        with http.session_transaction() as sesion:
            sesion['_user_id'] = str(usuario.id)
            sesion['_fresh'] = True
        datos = {'identificacion': '9001', 'fecha_gestion': '2026-01-15', 'medio': 'LLAMADA',
                 'contacto': 'Tesorería', 'observaciones': 'Promete abono.\nEnviar estado de cuenta.',
                 'fecha_compromiso': '2026-01-20', 'valor_compromiso': '50.25',
                 'proximo_seguimiento': '2026-01-21', 'usuario_nombre': 'Autor falso'}
        primero = http.post('/api/contable/seguimiento-cartera', json=datos)
        self.assertEqual(primero.status_code, 201)
        self.assertEqual(primero.json['seguimiento']['registrado_por'], 'Gestora cartera')
        self.assertTrue(primero.json['seguimiento']['fecha_hora_gestion'].startswith('2026-01-15T'))
        segundo = http.post('/api/contable/seguimiento-cartera', json=dict(
            datos, fecha_gestion='2026-01-16', observaciones='Confirma el compromiso.'))
        self.assertEqual(segundo.status_code, 201)
        db.session.remove()
        historial = http.get('/api/contable/seguimiento-cartera?identificacion=9001').json['seguimientos']
        self.assertEqual([item['fecha_gestion'] for item in historial], ['2026-01-16', '2026-01-15'])
        self.assertTrue(historial[1]['fecha_hora_gestion'].startswith('2026-01-15T'))
        self.assertEqual(historial[1]['valor_compromiso'], 50.25)
        self.assertEqual(historial[1]['observaciones'], datos['observaciones'])
        compromiso_id = primero.json['seguimiento']['id']
        g.pop('_login_user', None)
        ajeno = http.patch('/api/contable/seguimiento-cartera', json={'identificacion': '9002', 'id': compromiso_id})
        self.assertEqual(ajeno.status_code, 404)
        cumplido = http.patch('/api/contable/seguimiento-cartera', json={'identificacion': '9001', 'id': compromiso_id})
        self.assertEqual(cumplido.status_code, 200)
        self.assertEqual(cumplido.json['seguimiento']['estado_compromiso'], 'cumplido')
        self.assertEqual(cumplido.json['seguimiento']['compromiso_cumplido_por'], 'Gestora cartera')
        repetido = http.patch('/api/contable/seguimiento-cartera', json={'identificacion': '9001', 'id': compromiso_id})
        self.assertEqual(repetido.json, cumplido.json)
        with self.app.test_request_context(query_string={'informe': 'vencidas', 'fecha_corte': '2026-02-28'}):
            informe = contable.cartera_dinamica.__wrapped__().get_json()
        cliente = next(c for c in informe['clientes'] if c['identificacion'] == '9001')
        self.assertTrue(any(r['estado_compromiso'] == 'cumplido' for r in cliente['seguimientos']))
        self.assertEqual(http.get('/api/contable/seguimiento-cartera?identificacion=9002').json['seguimientos'], [])
        # Una variante del NIT con puntos y DV conserva el mismo historial.
        self.documento('FV', 3, '2026-01-01', [{'identificacion': '9.001-2', 'detalle': 'FV-2-3', 'debito': Decimal('10')}])
        self.assertEqual(len(http.get('/api/contable/seguimiento-cartera?identificacion=9.001-2').json['seguimientos']), 2)

    def test_seguimiento_agrupado_por_responsable_crea_registro_por_empresa(self):
        vendedor = Vendedor(nombre='Vendedora cartera')
        db.session.add(vendedor)
        db.session.flush()
        db.session.add_all([
            ClienteComercial(razon_social='Empresa grupo A', nit='9001', vendedor=vendedor,
                             responsable='Responsable Grupo', telefono_responsable='3001112222'),
            ClienteComercial(razon_social='Empresa grupo B', nit='9002', vendedor=vendedor,
                             responsable='Responsable Grupo', telefono_responsable='3001112222'),
        ])
        db.session.commit()
        def limpiar_maestro_grupo():
            db.session.query(ClienteComercial).filter(ClienteComercial.nit.in_(['9001', '9002'])).delete(synchronize_session=False)
            db.session.query(Vendedor).filter_by(nombre='Vendedora cartera').delete()
            db.session.commit()
        self.addCleanup(limpiar_maestro_grupo)
        self.factura_y_recibo()
        self.documento('FV', 2, '2026-01-01', [{'identificacion': '9002', 'nombre_tercero': 'Empresa grupo B',
                                                'detalle': 'FV-2-2', 'debito': Decimal('10')}])
        usuario = self.usuario_seguimiento()
        http = self.app.test_client()
        with http.session_transaction() as sesion:
            sesion['_user_id'] = str(usuario.id)
            sesion['_fresh'] = True
        respuesta = http.post('/api/contable/seguimiento-cartera', json={
            'identificacion': '9001',
            'alcance': 'grupo',
            'identificaciones_grupo': ['9001', '9002'],
            'fecha_gestion': '2026-01-15',
            'medio': ['CORREO', 'WHATSAPP'],
            'contacto': 'Responsable Grupo',
            'observaciones': 'Compromiso agrupado.',
            'fecha_compromiso': '2026-01-20',
            'valor_compromiso': '1.000,50',
        })
        self.assertEqual(respuesta.status_code, 201)
        self.assertEqual(len(respuesta.json['seguimientos']), 2)
        self.assertEqual(SiigoSeguimientoCartera.query.filter_by(identificacion='9001').count(), 1)
        self.assertEqual(SiigoSeguimientoCartera.query.filter_by(identificacion='9002').count(), 1)
        self.assertEqual(respuesta.json['seguimiento']['medio'], 'CORREO, WHATSAPP')
        self.assertEqual(respuesta.json['seguimiento']['valor_compromiso'], 1000.5)

    def test_seguimiento_valida_entradas_sin_guardar_registros_invalidos(self):
        self.factura_y_recibo()
        usuario = self.usuario_seguimiento()
        http = self.app.test_client()
        with http.session_transaction() as sesion:
            sesion['_user_id'] = str(usuario.id)
        datos = {'identificacion': '9001', 'fecha_gestion': '2026-01-15',
                 'medio': 'CORREO', 'observaciones': 'Solicita estado de cuenta.'}
        for cambios in [
            {'observaciones': '  '}, {'observaciones': 'x' * 5001}, {'medio': 'INVALIDO'},
            {'fecha_gestion': '2099-01-01'}, {'fecha_gestion': 'no-fecha'},
            {'fecha_compromiso': '2026-01-14'}, {'proximo_seguimiento': '2026-01-14'},
            {'valor_compromiso': '-1'}, {'valor_compromiso': 'NaN'}, {'valor_compromiso': 'Infinity'},
            {'valor_compromiso': '10'}, {'valor_compromiso': '1.123', 'fecha_compromiso': '2026-01-20'},
            {'identificacion': ''}, {'contacto': 'x' * 201},
        ]:
            with self.subTest(cambios=cambios):
                response = http.post('/api/contable/seguimiento-cartera', json=dict(datos, **cambios))
                self.assertEqual(response.status_code, 400)
        self.assertEqual(http.post('/api/contable/seguimiento-cartera', json=[]).status_code, 400)
        self.assertEqual(http.get('/api/contable/seguimiento-cartera?identificacion=INEXISTENTE').status_code, 404)
        self.assertEqual(SiigoSeguimientoCartera.query.count(), 0)

    def test_seguimiento_exige_sesion_y_permiso_de_ventas(self):
        self.permission.stop()
        http = self.app.test_client()
        self.assertEqual(http.get('/api/contable/seguimiento-cartera?identificacion=9001').status_code, 401)
        self.assertEqual(http.post('/api/contable/seguimiento-cartera', json={}).status_code, 401)
        usuario = self.usuario_seguimiento()
        with http.session_transaction() as sesion:
            sesion['_user_id'] = str(usuario.id)
        # setUp conserva un app_context; descartar el usuario anónimo de la consulta previa.
        g.pop('_login_user', None)
        self.assertEqual(http.get('/api/contable/seguimiento-cartera?identificacion=9001').status_code, 403)
        self.assertEqual(http.post('/api/contable/seguimiento-cartera', json={}).status_code, 403)

    def excel(self, ac_cuadrado=True):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['Comprobantes detallados'])
        sheet.append(['Secuencia', 'Fecha elaboracion', 'Codigo contable', 'Cuenta contable',
                      'Identificacion', 'Nombre tercero', 'Descripcion', 'Detalle', 'Debito', 'Credito'])
        for ref, debit, credit, other, description in [
            ('FV-2-1', 1000, 0, '41350501', 'Venta'),
            ('RC-1-1', 0, 900, '11100501', 'Abono FV-2-1'),
            ('AC-1-1', 0, 100, '13551501', 'Retenciones'),
        ]:
            sheet.append([f'Comprobante: {ref}'])
            sheet.append([1, '01/02/2026', '13050501', 'Clientes', '9001', 'Cliente prueba',
                          description, 'FV-2-1 Cuota: 1 Fecha: 31/01/2026', debit, credit])
            counter = credit if ac_cuadrado or not ref.startswith('AC') else 50
            sheet.append([2, '01/02/2026', other, 'Contrapartida', '9001', 'Cliente prueba',
                          description, None, counter, debit])
        content = BytesIO()
        workbook.save(content)
        return content.getvalue()

    def cargar(self, content):
        with self.app.test_request_context(method='POST', data={'archivo': (BytesIO(content), 'comprobantes.xlsx')}):
            with patch.object(contable, 'current_user', SimpleNamespace(id=None)):
                response = contable.cargar_comprobantes.__wrapped__()
        if isinstance(response, tuple):
            return response[0].get_json(), response[1]
        return response.get_json(), 200

    def test_reimportar_archivo_antiguo_completa_ac_sin_duplicados(self):
        content = self.excel()
        with patch.object(contable, '_comprobante_relevante', side_effect=lambda doc, cols: doc['tipo'] in {'FV', 'RC', 'NC', 'ND'}):
            data, status = self.cargar(content)
        self.assertEqual((status, data['comprobantes']), (200, 2))
        data, status = self.cargar(content)
        self.assertEqual((status, data['comprobantes']), (200, 1))
        self.assertEqual(SiigoCarga.query.count(), 1)
        carga = SiigoCarga.query.one()
        self.assertEqual((carga.registros_leidos, carga.registros_importados, carga.registros_omitidos), (3, 3, 0))
        self.assertEqual(carga.total_debito, Decimal('2000'))
        self.assertEqual(carga.total_credito, Decimal('2000'))
        data, status = self.cargar(content)
        self.assertEqual(status, 400)
        self.assertIn('no contiene comprobantes nuevos', data['error'])
        self.assertEqual(SiigoComprobante.query.count(), 3)
        self.assertEqual(SiigoMovimiento.query.count(), 6)
        self.assertEqual(self.cartera()['cartera_clientes'][0]['saldo'], 0)

    def test_cargue_comprobantes_crea_cliente_faltante_desde_excel(self):
        db.session.query(ClienteComercial).filter(ClienteComercial.nit.in_(['9001', '9001-2'])).delete(synchronize_session=False)
        db.session.commit()
        self.addCleanup(db.session.commit)
        self.addCleanup(lambda: db.session.query(ClienteComercial).filter(ClienteComercial.nit.in_(['9001', '9001-2'])).delete(synchronize_session=False))
        self.assertIsNone(ClienteComercial.query.filter_by(nit='9001').first())
        data, status = self.cargar(self.excel())
        self.assertEqual(status, 200)
        self.assertEqual(data['clientes_creados'], 1)
        cliente = ClienteComercial.query.filter_by(nit='9001').one()
        self.assertEqual(cliente.razon_social, 'Cliente prueba')
        self.assertTrue(cliente.importado_siigo)

    def test_cargue_comprobantes_no_duplica_cliente_existente_por_identificacion(self):
        db.session.query(ClienteComercial).filter(ClienteComercial.nit.in_(['9001', '9001-2'])).delete(synchronize_session=False)
        db.session.commit()
        self.addCleanup(db.session.commit)
        self.addCleanup(lambda: db.session.query(ClienteComercial).filter(ClienteComercial.nit.in_(['9001', '9001-2'])).delete(synchronize_session=False))
        db.session.add(ClienteComercial(nit='9001-2', razon_social='Cliente existente'))
        db.session.commit()
        data, status = self.cargar(self.excel())
        self.assertEqual(status, 200)
        self.assertEqual(data['clientes_creados'], 0)
        self.assertEqual(ClienteComercial.query.filter(ClienteComercial.nit.in_(['9001', '9001-2'])).count(), 1)

    def test_ac_descuadrado_revierte_toda_la_importacion(self):
        data, status = self.cargar(self.excel(ac_cuadrado=False))
        self.assertEqual(status, 400)
        self.assertIn('AC-1-1 no cuadra', data['error'])
        self.assertEqual(SiigoCarga.query.count(), 0)
        self.assertEqual(SiigoComprobante.query.count(), 0)
        self.assertEqual(SiigoMovimiento.query.count(), 0)

    def test_todos_los_tipos_cruzan_con_la_factura_sin_duplicar_contrapartidas(self):
        self.documento('FV', 1, '2026-01-01', [{'debito': Decimal('1000')}])
        self.documento('CC', 1, '2026-02-01', [
            {'credito': Decimal('100')}, {'debito': Decimal('30')},
            {'codigo_contable': '13990501', 'debito': Decimal('70')},
        ])
        self.documento('DF', 1, '2026-02-02', [{'credito': Decimal('200')}])
        self.documento('XYZ', 1, '2026-02-03', [{'credito': Decimal('730')}])
        self.documento('EG', 1, '2026-02-04', [{'debito': Decimal('20')}])
        for corte, saldo in [('2026-02-01', 930), ('2026-02-02', 730),
                             ('2026-02-03', 0), ('2026-02-04', 20)]:
            with self.subTest(corte=corte):
                cartera = self.cartera(corte)
                factura = cartera['cartera_clientes'][0]['facturas'][0]
                self.assertEqual(factura['saldo'], saldo)
                self.assertEqual(factura['otros_movimientos'], 1000 - saldo)
                self.assertEqual(sum(m['debito'] - m['credito'] for m in factura['movimientos']), saldo)
                self.assertEqual(cartera['pagos_clientes'], [])
                with self.app.test_request_context(query_string={'informe': 'vencidas', 'fecha_corte': corte}):
                    vencidas = contable.cartera_dinamica.__wrapped__().get_json()
                self.assertEqual(vencidas['cantidad_facturas'], int(saldo > 0))

    def test_una_fv_que_cruza_otra_no_se_suma_dos_veces(self):
        self.documento('FV', 1, '2026-01-01', [{'debito': Decimal('1000')}])
        self.documento('FV', 2, '2026-02-01', [
            {'debito': Decimal('500'), 'detalle': 'FV-2-2 Fecha: 15/02/2026'},
            {'credito': Decimal('200'), 'detalle': 'FV-2-1'},
        ])
        data = self.cartera()
        facturas = {f['referencia']: f for c in data['cartera_clientes'] for f in c['facturas']}
        self.assertEqual(facturas['FV-2-1']['saldo'], 800)
        self.assertEqual(facturas['FV-2-1']['facturado'], 1000)
        self.assertEqual(facturas['FV-2-1']['otros_movimientos'], 200)
        self.assertEqual(facturas['FV-2-2']['saldo'], 500)
        self.assertEqual(data['cartera_clientes'][0]['saldo'], 1300)

    def test_cargue_recupera_cualquier_tipo_con_cartera_y_exporta_trazabilidad(self):
        libro = load_workbook(BytesIO(self.excel()))
        for fila in libro.active:
            if fila[0].value == 'Comprobante: AC-1-1':
                fila[0].value = 'Comprobante: NUEVO-9-1'
        archivo = BytesIO()
        libro.save(archivo)
        libro.close()
        with patch.object(contable, '_comprobante_relevante', side_effect=lambda doc, cols: doc['tipo'] in {'FV', 'RC'}):
            data, status = self.cargar(archivo.getvalue())
        self.assertEqual((status, data['comprobantes']), (200, 2))
        data, status = self.cargar(archivo.getvalue())
        self.assertEqual((status, data['comprobantes']), (200, 1))
        self.assertEqual(self.cartera()['cartera_clientes'][0]['saldo'], 0)
        self.assertEqual(SiigoComprobante.query.filter_by(tipo_documento='NUEVO').count(), 1)
        self.assertEqual(self.cargar(archivo.getvalue())[1], 400)
        self.documento('REV', 2, '2026-02-02', [{'debito': Decimal('25')}])
        with self.app.test_request_context(query_string={'informe': 'vencidas', 'formato': 'xlsx', 'fecha_corte': '2026-02-28'}):
            response = contable.cartera_dinamica.__wrapped__()
            response.direct_passthrough = False
            libro = load_workbook(BytesIO(response.get_data()))
            self.assertEqual(libro['Facturas vencidas']['G4'].value, 25)
            filas = list(libro['Cruces por factura'].values)[1:]
            self.assertEqual(sum(f[5] - f[6] for f in filas), 25)
            self.assertEqual({f[2] for f in filas}, {'FV-2-1', 'RC-1-1', 'NUEVO-9-1', 'REV-1-2'})
            libro.close()
            response.close()


if __name__ == '__main__':
    unittest.main()
