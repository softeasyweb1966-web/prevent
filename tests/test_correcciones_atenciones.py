"""Flujo real Excel -> revision -> actualizacion atomica -> auditoria -> regeneracion."""
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from unittest import TestCase
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook, Workbook

import test_prefacturas_empresa as base
from app import correcciones_atenciones as correcciones
from app.models import db, AtencionDiaDetalle, AuditLog, PrefacturaComercial
from app.routes import cargue_atenciones as ca


class CorreccionesTest(TestCase):
    def setUp(self):
        self.f = base.PrefacturasEmpresaTest()
        self.f.setUp()
        self.f.app.add_url_rule('/corregir', view_func=correcciones.cargar_correcciones_excel.__wrapped__, methods=['POST'])
        self.f.app.add_url_rule('/informe', view_func=correcciones.informe_correcciones.__wrapped__)
        self.f.app.add_url_rule('/regenerar', view_func=ca.regenerar_prefactura_empresa.__wrapped__)
        self.patches = [patch.object(ca, '_desvincular_prefacturas_de_atencion'),
                        patch.object(ca, '_sincronizar_cruce_prefacturas_con_atencion')]
        for p in self.patches:
            p.start()
        ca.current_user.usuario = 'operador.prueba'
        response = self.f.generar()
        self.assertEqual(response.status_code, 200)
        self.archivos = [self.excel_original(i) for i in (1, 2)]

    def excel_original(self, cliente_id):
        wb = Workbook(); ws = wb.active
        ws.append(list(ca.COLUMNAS_ESPERADAS_CANONICAS))
        for reg in AtencionDiaDetalle.query.filter_by(cliente_id=cliente_id).order_by(AtencionDiaDetalle.id):
            ws.append([getattr(reg, campo) for campo in correcciones.CAMPOS_EXCEL])
        buf = BytesIO(); wb.save(buf)
        return f'Atenciones-{cliente_id}.xlsx', buf.getvalue()

    def tearDown(self):
        for p in reversed(self.patches):
            p.stop()
        self.f.tearDown()

    def modificar(self, index=0, cambios=None):
        nombre, contenido = self.archivos[index]
        wb = load_workbook(BytesIO(contenido))
        ws = wb.active
        for celda, valor in (cambios or {'H2': 'Paciente corregido', 'D2': 350}).items():
            ws[celda] = valor
        buf = BytesIO()
        wb.save(buf)
        wb.close()
        return nombre, buf.getvalue()

    def cargar(self, archivos, token=None):
        data = {'periodo_desde': '2026-09-01', 'periodo_hasta': '2026-09-15', 'archivos': [(BytesIO(contenido), nombre) for nombre, contenido in archivos]}
        if token:
            data.update(accion='aplicar', token=token)
        return self.f.http.post('/corregir', data=data, content_type='multipart/form-data')

    def test_revisar_aplicar_informe_y_regenerar(self):
        archivos = [self.modificar(), self.modificar(1, {'G2': '009999', 'H2': 'Paciente dos corregido'})]
        response = self.cargar(archivos)
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(response.json['atenciones'], 2)
        self.assertEqual(AuditLog.query.count(), 0)
        self.assertEqual(db.session.get(AtencionDiaDetalle, 1).nombre_paciente, 'Paciente 1')
        response = self.cargar(archivos, response.json['token'])
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(response.json['actualizadas'], 2)
        logs = AuditLog.query.order_by(AuditLog.id).all()
        self.assertEqual(len(logs), 2)
        self.assertEqual(logs[0].datos_anteriores['nombre_paciente'], 'Paciente 1')
        self.assertEqual(logs[0].datos_nuevos['nombre_paciente'], 'Paciente corregido')
        self.assertEqual(logs[0].usuario_id, 1)
        self.assertIsNotNone(logs[0].created_at)
        self.assertEqual(PrefacturaComercial.query.filter_by(cliente_id=1).one().valor_total, Decimal('350'))
        informe = self.f.http.get('/informe')
        self.assertEqual(informe.status_code, 200, informe.json)
        self.assertEqual(informe.json['total'], 2)
        self.assertEqual(len(informe.json['filas']), 4)
        self.assertEqual(informe.json['filas'][0]['usuario'], 'operador.prueba')
        excel = self.f.http.get('/informe?formato=xlsx')
        self.assertEqual(load_workbook(BytesIO(excel.data)).active.max_row, 5)
        regenerado = self.f.http.get('/regenerar?empresa=Empresa+1&fecha_desde=2026-09-01&fecha_hasta=2026-09-15')
        self.assertEqual(regenerado.status_code, 200)
        with ZipFile(BytesIO(regenerado.data)) as z:
            wb = load_workbook(BytesIO(z.read(z.namelist()[0])))
            self.assertEqual(wb['relacion-pacientes']['C4'].value, 'Paciente corregido')
            self.assertEqual(wb['relacion-pacientes']['E4'].value, 350)
            self.assertNotIn('corregir-atenciones', wb.sheetnames)

    def test_excel_sin_cambios(self):
        response = self.cargar(self.archivos)
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(response.json['atenciones'], 0)

    def test_lote_invalido_no_modifica_ninguna_atencion(self):
        response = self.cargar([self.modificar(), self.modificar(1, {'D2': -1})])
        self.assertEqual(response.status_code, 400)
        self.assertIn('fila 2', response.json['error'])
        self.assertEqual(AuditLog.query.count(), 0)
        self.assertEqual(db.session.get(AtencionDiaDetalle, 1).precio, Decimal(100))

    def test_rechaza_archivo_obsoleto(self):
        archivo = self.modificar()
        preview = self.cargar([archivo]).json
        reg = db.session.get(AtencionDiaDetalle, 1)
        reg.precio = 999
        db.session.commit()
        response = self.cargar([archivo], preview['token'])
        self.assertEqual(response.status_code, 400)
        self.assertIn('cambiaron', response.json['error'])
        self.assertEqual(AuditLog.query.count(), 0)

    def test_rechaza_id_alterado_formula_y_fuera_de_periodo(self):
        for cambios in ({'A2': '999'}, {'D2': '=100+50'}, {'M2': datetime(2026, 10, 1)}, {'D2': -3}):
            response = self.cargar([self.modificar(cambios=cambios)])
            self.assertEqual(response.status_code, 400, response.json)
        self.assertEqual(AuditLog.query.count(), 0)

    def test_rechaza_duplicados_y_archivo_sin_hoja(self):
        archivo = self.modificar()
        self.assertEqual(self.cargar([archivo, archivo]).status_code, 400)
        prefactura = self.f.generar('&cliente_id=1')
        with ZipFile(BytesIO(prefactura.data)) as z:
            nombre = next(n for n in z.namelist() if n.startswith('cred-'))
            response = self.cargar([(nombre, z.read(nombre))])
        self.assertEqual(response.status_code, 400)
        self.assertIn('Columnas requeridas', response.json['error'])

    def test_rechaza_cambios_no_revisados(self):
        token = self.cargar([self.modificar()]).json['token']
        response = self.cargar([self.modificar(cambios={'D2': 500})], token)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(AuditLog.query.count(), 0)

    def test_no_modifica_prefacturas_cerradas_o_bloqueadas(self):
        pref = PrefacturaComercial.query.filter_by(cliente_id=1).one()
        pref.estado = 'CERRADA'; db.session.commit()
        response = self.cargar([self.modificar()])
        self.assertEqual(response.status_code, 400)
        pref.estado = 'BORRADOR'; pref.bloqueada_por_pago = True; db.session.commit()
        self.assertEqual(self.cargar([self.modificar()]).status_code, 400)

    def test_fallo_guardado_revierte_atenciones_y_auditoria(self):
        archivo = self.modificar()
        token = self.cargar([archivo]).json['token']
        with patch.object(correcciones, '_actualizar_borradores', side_effect=RuntimeError('fallo simulado')):
            self.assertEqual(self.cargar([archivo], token).status_code, 500)
        self.assertEqual(db.session.get(AtencionDiaDetalle, 1).precio, Decimal(100))
        self.assertEqual(AuditLog.query.count(), 0)

    def test_anulacion_y_cambio_forma_pago_actualizan_totales(self):
        archivos = [self.modificar(cambios={'O2': 'ANULADA'}), self.modificar(1, {'E2': 'EFECTIVO'})]
        token = self.cargar(archivos).json['token']
        response = self.cargar(archivos, token)
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(PrefacturaComercial.query.filter_by(cliente_id=1).one().valor_total, 0)
        self.assertEqual(PrefacturaComercial.query.filter_by(cliente_id=2).one().forma_pago, 'EFECTIVO')
        self.assertIsNotNone(db.session.get(AtencionDiaDetalle, 1).fecha_anulacion)

    def test_rechaza_sin_permiso(self):
        with patch.object(ca, '_asegurar_acceso_registro_atencion', side_effect=PermissionError('Sin acceso')):
            self.assertEqual(self.cargar([self.modificar()]).status_code, 403)
        self.assertEqual(AuditLog.query.count(), 0)

    def test_auditoria_permanece_si_se_elimina_atencion(self):
        archivo = self.modificar()
        token = self.cargar([archivo]).json['token']
        self.assertEqual(self.cargar([archivo], token).status_code, 200)
        db.session.delete(db.session.get(AtencionDiaDetalle, 1)); db.session.commit()
        self.assertEqual(self.f.http.get('/informe').json['total'], 1)

    def test_examenes_agrupados_se_corrigen_individualmente(self):
        reg = db.session.get(AtencionDiaDetalle, 1)
        db.session.add(AtencionDiaDetalle(
            cargue_id=reg.cargue_id, cliente_id=reg.cliente_id, nro_orden=reg.nro_orden,
            nro_identificacion=reg.nro_identificacion, nombre_paciente=reg.nombre_paciente,
            fecha_creacion_orden=reg.fecha_creacion_orden, servicio='Audiometria',
            precio=50, forma_pago='CREDITO', estado_orden='ACTIVA', archivo_origen='test.xlsx'))
        db.session.commit()
        response = self.f.generar('&cliente_id=1')
        self.archivos = [self.excel_original(1)]
        archivo = self.modificar(cambios={'D2': 120})
        token = self.cargar([archivo]).json['token']
        self.assertEqual(self.cargar([archivo], token).status_code, 200)
        self.assertEqual(db.session.get(AtencionDiaDetalle, 1).precio, Decimal(120))
        self.assertEqual(db.session.get(AtencionDiaDetalle, 3).precio, Decimal(50))
        self.assertEqual(PrefacturaComercial.query.filter_by(cliente_id=1).one().valor_total, Decimal(170))

    def test_cierre_entre_revision_y_aplicacion(self):
        archivo = self.modificar()
        token = self.cargar([archivo]).json['token']
        pref = PrefacturaComercial.query.filter_by(cliente_id=1).one()
        pref.estado = 'CERRADA'; db.session.commit()
        self.assertEqual(self.cargar([archivo], token).status_code, 400)
        self.assertEqual(AuditLog.query.count(), 0)

    def test_informe_filtra_cartera_del_vendedor(self):
        archivos = [self.modificar(), self.modificar(1)]
        token = self.cargar(archivos).json['token']
        self.assertEqual(self.cargar(archivos, token).status_code, 200)
        from unittest.mock import Mock
        from app.models import ClienteComercial
        with patch.object(ca, '_is_admin_user', return_value=False), \
             patch.object(ca, '_resolver_vendedor_usuario_actual', return_value=Mock(id=1)), \
             patch.object(ca, '_clientes_visibles_query', return_value=ClienteComercial.query.filter_by(id=1)):
            response = self.f.http.get('/informe')
            self.assertEqual(response.status_code, 200, response.json)
            self.assertEqual(response.json['total'], 1)
            self.assertTrue(all(f['empresa'] == 'Empresa 1' for f in response.json['filas']))

    def test_cambio_empresa_actualiza_ambas_prefacturas_y_auditoria(self):
        archivo = self.modificar(cambios={'I2': 'Empresa 2'})
        preview = self.cargar([archivo])
        self.assertEqual(preview.status_code, 200, preview.json)
        response = self.cargar([archivo], preview.json['token'])
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(db.session.get(AtencionDiaDetalle, 1).cliente_id, 2)
        self.assertEqual(PrefacturaComercial.query.filter_by(cliente_id=1).one().valor_total, 0)
        self.assertEqual(PrefacturaComercial.query.filter_by(cliente_id=2).one().valor_total, 300)
        self.assertEqual(len(response.json['periodos']), 2)
        log = AuditLog.query.one()
        self.assertEqual(log.datos_anteriores['cliente_id'], 1)
        self.assertEqual(log.datos_nuevos['cliente_id'], 2)

    def test_coincidencia_ambigua_no_modifica_ninguna_fila(self):
        reg = db.session.get(AtencionDiaDetalle, 1)
        db.session.add(AtencionDiaDetalle(cargue_id=reg.cargue_id, cliente_id=1,
            nro_orden=reg.nro_orden, servicio=reg.servicio, nro_identificacion=reg.nro_identificacion,
            fecha_creacion_orden=reg.fecha_creacion_orden, precio=80, estado_gestion='CARGADA'))
        db.session.commit()
        response = self.cargar([self.modificar()])
        self.assertEqual(response.status_code, 400)
        self.assertIn('atencion unica', response.json['error'])
        self.assertEqual(AuditLog.query.count(), 0)

    def test_factura_sede_y_fechas_del_formato_original_se_auditan(self):
        archivo = self.modificar(cambios={'B2': 'FAC-123', 'C2': datetime(2026, 9, 5), 'K2': 'Sede Norte'})
        token = self.cargar([archivo]).json['token']
        self.assertEqual(self.cargar([archivo], token).status_code, 200)
        self.assertEqual(db.session.get(AtencionDiaDetalle, 1).nro_factura, 'FAC-123')
        self.assertEqual(len(self.f.http.get('/informe').json['filas']), 3)
        self.assertEqual(self.cargar([archivo]).json['atenciones'], 0)

    def test_repetir_excel_anulado_no_borra_fecha_ni_duplica_auditoria(self):
        archivo = self.modificar(cambios={'O2': 'ANULADA'})
        token = self.cargar([archivo]).json['token']
        self.assertEqual(self.cargar([archivo], token).status_code, 200)
        self.assertEqual(self.cargar([archivo]).json['atenciones'], 0)
        self.assertEqual(AuditLog.query.count(), 1)
