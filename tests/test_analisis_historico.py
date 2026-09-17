"""Excel de análisis sobre PostgreSQL aislado; no se modifica la base operativa."""
from datetime import datetime
from decimal import Decimal
from io import BytesIO
import hashlib
import hmac
import json
import unittest

from openpyxl import load_workbook
from sqlalchemy import event

import test_clientes_maestro as maestro
from app.models import (db, AtencionDiaDetalle, CargueAtencionDia,
                        ComercialCatalogoItem, ComercialPaqueteDetalle,
                        ClienteComercialTarifa)
from app.analisis_historico import manifiesto_hojas


class AnalisisHistoricoTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        maestro.MaestroTest.setUpClass()

    def setUp(self):
        self.f = maestro.MaestroTest()
        self.f.setUp(); self.f.login(0)
        self.a = ComercialCatalogoItem(nombre='Consulta altura', tipo_item='EXAMEN', tipo_examen='CONSULTA', tarifa_base=100)
        self.b = ComercialCatalogoItem(nombre='Audiometría', tipo_item='EXAMEN', tipo_examen='PARACLINICO', tarifa_base=200)
        self.p = ComercialCatalogoItem(nombre='Alturas existente', tipo_item='PAQUETE', tarifa_base=900)
        self.carga = CargueAtencionDia(nombre_archivo='historico.xlsx')
        db.session.add_all([self.a, self.b, self.p, self.carga]); db.session.flush()
        db.session.add_all([ComercialPaqueteDetalle(paquete_id=self.p.id, examen_id=i.id, cantidad=1) for i in (self.a,self.b)])
        self.tarifa = ClienteComercialTarifa(cliente_id=self.f.c1.id, catalogo_item_id=self.p.id,
            tarifa_negociada=Decimal('99000'), activo=False, vigencia_hasta=datetime(2025,12,31))
        db.session.add(self.tarifa)
        for numero, fecha in enumerate((datetime(2026,1,1), datetime(2026,1,20), datetime(2026,2,28,23,59,59,999999)),1):
            self.orden(self.f.c1, str(numero), fecha, Decimal('40000'))
        self.orden(self.f.c2, 'otra', datetime(2026,2,1), Decimal('46000'))
        db.session.commit()

    def tearDown(self):
        self.f.tearDown()

    def orden(self, cliente, numero, fecha, precio):
        for item in (self.a, self.b):
            db.session.add(AtencionDiaDetalle(cargue_id=self.carga.id, cliente_id=cliente.id,
                vendedor_id=cliente.vendedor_id, nro_orden=numero, nro_identificacion='pac-'+numero,
                fecha_creacion_orden=fecha, servicio=item.nombre, precio=precio,
                forma_pago='CRÉDITO', estado_orden='ACTIVA', archivo_origen='historico.xlsx'))

    def descargar(self, extra=''):
        return self.f.http.get('/api/comercial/tarifas/analisis-historico.xlsx?fecha_desde=2026-01-01&fecha_hasta=2026-02-28'+extra)

    @staticmethod
    def filas(wb, hoja):
        rows = list(wb[hoja].values)
        return [dict(zip(rows[0], row)) for row in rows[1:]]

    def test_compartir_composicion_conservar_tarifas_y_cero_escrituras(self):
        sentencias = []
        def capturar(conn, cursor, statement, parameters, context, executemany):
            sentencias.append(statement.strip().split()[0].upper())
        event.listen(db.engine, 'before_cursor_execute', capturar)
        try:
            r = self.descargar()
        finally:
            event.remove(db.engine, 'before_cursor_execute', capturar)
        self.assertEqual(r.status_code, 200, r.json)
        self.assertFalse({'INSERT','UPDATE','DELETE','ALTER','CREATE'} & set(sentencias))
        wb = load_workbook(BytesIO(r.data))
        paquetes = self.filas(wb, 'PAQUETES')
        self.assertEqual(len(paquetes), 1)
        self.assertIsNone(paquetes[0]['Nombre paquete'])
        self.assertEqual(paquetes[0]['Cantidad empresas'], 2)
        self.assertIn('Alturas existente', paquetes[0]['Paquetes existentes'])
        relaciones = self.filas(wb, 'EMPRESA_PAQUETE')
        self.assertEqual({r['ID composición'] for r in relaciones}, {paquetes[0]['ID composición']})
        self.assertEqual({r['Provisional'] for r in relaciones}, {'Paquete 1'})
        self.assertEqual(relaciones[0]['Cantidad usos'], 3)
        self.assertEqual(relaciones[0]['Clasificación'], 'Paquete provisional recurrente')
        self.assertEqual(relaciones[1]['Clasificación'], 'Combinación ocasional')
        self.assertIn('Ya existe', relaciones[0]['Propuesta'])
        self.assertIn('99000', relaciones[0]['Relaciones existentes'])
        self.assertEqual({r['Valor histórico'] for r in self.filas(wb,'VALORES_PAQUETE')}, {80000,92000})
        self.assertEqual(self.tarifa.tarifa_negociada, Decimal('99000'))
        self.assertFalse(self.tarifa.activo)
        self.assertEqual(len(self.filas(wb, 'EVIDENCIA')),8)

    def test_formato_firmado_admite_nombre_pero_detecta_cambio_composicion(self):
        r = self.descargar()
        wb = load_workbook(BytesIO(r.data))
        control = wb['_CONTROL']
        raw = ''.join(row[1] for row in list(control.values)[3:])
        firma = hmac.new(b'test-only', raw.encode('utf-8'), hashlib.sha256).hexdigest()
        self.assertEqual(control['B3'].value, firma)
        original = json.loads(raw)['hojas']
        self.assertEqual(manifiesto_hojas(wb), original)
        wb['PAQUETES']['C2'] = 'Paquete Alturas Básico'
        self.assertFalse(wb['PAQUETES']['C2'].protection.locked)
        self.assertTrue(wb['COMPONENTES']['E2'].protection.locked)
        buf = BytesIO(); wb.save(buf); buf.seek(0)
        editado = load_workbook(buf)
        self.assertEqual(manifiesto_hojas(editado), original)
        editado['COMPONENTES']['E2'] = 2
        self.assertNotEqual(manifiesto_hojas(editado), original)

    def test_alcance_vendedor_filtro_cliente_y_permiso(self):
        self.f.login(1)
        r = self.descargar()
        self.assertEqual(r.status_code,200)
        wb = load_workbook(BytesIO(r.data))
        self.assertEqual({r['ID empresa'] for r in self.filas(wb,'EMPRESA_PAQUETE')}, {self.f.c1.id})
        self.assertEqual(self.descargar(f'&cliente_id={self.f.c2.id}').status_code,403)
        self.f.login(0)
        wb = load_workbook(BytesIO(self.descargar(f'&cliente_id={self.f.c2.id}').data))
        self.assertEqual({r['ID empresa'] for r in self.filas(wb,'EMPRESA_PAQUETE')}, {self.f.c2.id})
        self.f.users[1].role.permisos = [p for p in self.f.users[1].role.permisos if p.nombre != 'comercial_atenciones_read']
        db.session.commit(); self.f.login(1)
        self.assertEqual(self.descargar().status_code,403)

    def test_incidencias_valores_y_estabilidad(self):
        original = load_workbook(BytesIO(self.descargar().data))
        self.orden(self.f.c1,'cambio',datetime(2026,2,10),Decimal('41000'))
        self.orden(self.f.c1,'nulo',datetime(2026,2,11),None)
        self.orden(self.f.c1,'duplicada',datetime(2026,2,12),Decimal('1'))
        self.orden(self.f.c1,'duplicada',datetime(2026,2,12),Decimal('1'))
        db.session.add(AtencionDiaDetalle(cargue_id=self.carga.id, cliente_id=self.f.c1.id,
            nro_orden='sin-catalogo', nro_identificacion='x',fecha_creacion_orden=datetime(2026,2,2),
            servicio='=NO_EJECUTAR()',precio=1, forma_pago='EFECTIVO'))
        db.session.commit()
        wb = load_workbook(BytesIO(self.descargar().data))
        self.assertEqual(wb['PAQUETES']['A2'].value,original['PAQUETES']['A2'].value)
        relaciones = self.filas(wb,'EMPRESA_PAQUETE')
        self.assertIn('Múltiples valores', relaciones[0]['Observaciones'])
        self.assertIn('Precio nulo', relaciones[0]['Observaciones'])
        motivos = [r['Motivo'] for r in self.filas(wb,'INCIDENCIAS')]
        self.assertTrue(any('repetido' in x for x in motivos))
        self.assertTrue(any('coincidencia' in x for x in motivos))
        celdas = [c for row in wb['INCIDENCIAS'] for c in row if c.value == '=NO_EJECUTAR()']
        self.assertTrue(celdas)
        self.assertTrue(all(c.data_type=='s' for c in celdas))

    def test_fechas_invalidas(self):
        self.assertEqual(self.f.http.get('/api/comercial/tarifas/analisis-historico.xlsx').status_code,400)
        self.assertEqual(self.f.http.get('/api/comercial/tarifas/analisis-historico.xlsx?fecha_desde=2026-03-01&fecha_hasta=2026-01-01').status_code,400)

    def test_empresa_origen_se_analiza_sin_vincular_historico(self):
        self.orden(self.f.c1, 'origen', datetime(2026,2,15),Decimal('39000'))
        db.session.flush()
        registros = AtencionDiaDetalle.query.filter_by(nro_orden='origen').all()
        for r in registros:
            r.cliente_id = None
            r.acuerdo_comercial = self.f.c1.razon_social
        db.session.commit(); self.f.login(1)
        wb = load_workbook(BytesIO(self.descargar(f'&cliente_id={self.f.c1.id}').data))
        origen = self.filas(wb,'EMPRESAS_ORIGEN')[0]
        self.assertEqual(origen['IDs clientes candidatos'],str(self.f.c1.id))
        self.assertTrue(origen['ID empresa análisis'].startswith('ORIG-'))
        propuestas = self.filas(wb,'EMPRESA_PAQUETE')
        self.assertTrue(any(r['Propuesta'].startswith('Pendiente identidad') for r in propuestas))
        self.assertTrue(all(r.cliente_id is None for r in registros))
