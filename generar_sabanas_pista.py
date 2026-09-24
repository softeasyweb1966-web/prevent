"""Carga PISTA.xlsx y genera las sabanas de prefacturas.

Uso:
    .venv\Scripts\python.exe generar_sabanas_pista.py
"""

from __future__ import annotations

import os
import io
import re
import zipfile
from difflib import SequenceMatcher
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch


BASE_DIR = Path(__file__).resolve().parent
ARCHIVO_PISTA = BASE_DIR / "PISTA.xlsx"
CARPETA_SALIDA = BASE_DIR / "sabanas"
FECHA_DESDE = datetime(2024, 1, 1)
FECHA_HASTA = datetime.now().replace(hour=23, minute=59, second=59, microsecond=0)
DATABASE_URL_LOCAL = "postgresql+psycopg2://postgres:PreventPg2026Local1@127.0.0.1:5432/prevent_utf8"
PALABRAS_RUIDO_EMPRESA = {
    "s",
    "sa",
    "sas",
    "ltda",
    "limitada",
    "cia",
    "compania",
    "compañia",
    "empresa",
    "grupo",
    "de",
    "del",
    "la",
    "las",
    "los",
    "y",
    "en",
    "para",
    "con",
}

os.environ.setdefault("DATABASE_URL", DATABASE_URL_LOCAL)
os.environ.setdefault("FLASK_ENV", "production")

from app import create_app  # noqa: E402
from app.models import Usuario  # noqa: E402
from app.routes import cargue_atenciones as ca  # noqa: E402


def _usuario_para_proceso() -> SimpleNamespace:
    usuario = (
        Usuario.query.filter_by(is_easy=True, activo=True).first()
        or Usuario.query.filter_by(usuario="admin", activo=True).first()
        or Usuario.query.filter_by(activo=True).order_by(Usuario.id.asc()).first()
    )
    return SimpleNamespace(id=usuario.id if usuario else None)


def leer_empresas_pista() -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(ARCHIVO_PISTA, read_only=True, data_only=True)
    ws = wb.active
    empresas = []
    vistas = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            normalizada = ca._normalizar_match(str(row[0]).strip())
            if normalizada and normalizada not in vistas:
                empresas.append(normalizada)
                vistas.add(normalizada)
    return empresas


def _normalizar_nombre_archivo_sabana(nombre_archivo: str) -> str:
    stem = Path(nombre_archivo).stem
    stem = re.sub(r"^(cred|efec|mixto)-", "", stem, flags=re.IGNORECASE)
    stem = re.sub(r"-\d{8}-\d{8}$", "", stem)
    return ca._normalizar_match(stem.replace("_", " "))


def _tokens_empresa(nombre: str) -> set[str]:
    return {
        token
        for token in re.split(r"\s+", ca._normalizar_match(nombre))
        if len(token) >= 3 and token not in PALABRAS_RUIDO_EMPRESA
    }


def _coincide_empresa(nombre_sabana: str, empresas_pista: list[str]) -> tuple[bool, str, str]:
    tokens_sabana = _tokens_empresa(nombre_sabana)
    mejor_empresa = ""
    mejor_motivo = ""
    mejor_puntaje = 0.0

    for empresa in empresas_pista:
        if nombre_sabana == empresa:
            return True, empresa, "exacta"

        if nombre_sabana in empresa or empresa in nombre_sabana:
            return True, empresa, "contenida"

        tokens_pista = _tokens_empresa(empresa)
        comunes = tokens_sabana & tokens_pista
        if comunes:
            cobertura_pista = len(comunes) / max(len(tokens_pista), 1)
            cobertura_sabana = len(comunes) / max(len(tokens_sabana), 1)
            puntaje_tokens = min(cobertura_pista, cobertura_sabana)
            if puntaje_tokens > mejor_puntaje:
                mejor_puntaje = puntaje_tokens
                mejor_empresa = empresa
                mejor_motivo = f"palabras {puntaje_tokens:.0%}: {', '.join(sorted(comunes))}"

        puntaje_texto = SequenceMatcher(None, nombre_sabana, empresa).ratio()
        if puntaje_texto > mejor_puntaje:
            mejor_puntaje = puntaje_texto
            mejor_empresa = empresa
            mejor_motivo = f"similitud {puntaje_texto:.0%}"

    if mejor_puntaje >= 0.82:
        return True, mejor_empresa, mejor_motivo
    return False, mejor_empresa, mejor_motivo


def _guardar_reporte_coincidencias(reporte: list[dict]) -> Path:
    from openpyxl import Workbook

    salida = CARPETA_SALIDA / "Reporte-coincidencias-PISTA.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "coincidencias"
    ws.append(["incluida", "archivo_sabana", "empresa_sabana", "empresa_pista", "motivo"])
    for fila in reporte:
        ws.append([
            "SI" if fila["incluida"] else "NO",
            fila["archivo_sabana"],
            fila["empresa_sabana"],
            fila["empresa_pista"],
            fila["motivo"],
        ])
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 65
    ws.column_dimensions["C"].width = 55
    ws.column_dimensions["D"].width = 55
    ws.column_dimensions["E"].width = 45
    wb.save(salida)
    return salida


def generar_zip_sabanas(app, usuario, empresas_pista: list[str]) -> tuple[Path, Path, int, int]:
    periodo = f"{FECHA_DESDE:%d%m%Y}-{FECHA_HASTA:%d%m%Y}"
    salida = CARPETA_SALIDA / f"Prefacturas-{periodo}.zip"
    query = f"/api/comercial/prefacturas/generar?fecha_desde={FECHA_DESDE:%Y-%m-%d}&fecha_hasta={FECHA_HASTA:%Y-%m-%d}"

    with app.test_request_context(query):
        with patch.object(ca, "_require_commercial_permission"), \
             patch.object(ca, "_is_admin_user", return_value=True), \
             patch.object(ca, "_resolver_vendedor_usuario_actual", return_value=None), \
             patch.object(ca, "current_user", usuario), \
             patch.object(ca, "exigir_cliente"):
            respuesta = ca.generar_prefacturas.__wrapped__()

    if getattr(respuesta, "status_code", 200) != 200:
        detalle = respuesta.get_json(silent=True) if hasattr(respuesta, "get_json") else None
        raise RuntimeError(f"No se pudieron generar las sabanas: {detalle or respuesta}")

    respuesta.direct_passthrough = False
    CARPETA_SALIDA.mkdir(exist_ok=True)
    total_generadas = 0
    total_filtradas = 0
    with zipfile.ZipFile(io.BytesIO(respuesta.get_data()), "r") as zin:
        with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as zout:
            reporte = []
            for item in zin.infolist():
                if item.filename == "resumen_periodo.xlsx":
                    continue
                total_generadas += 1
                empresa_sabana = _normalizar_nombre_archivo_sabana(item.filename)
                incluida, empresa_pista, motivo = _coincide_empresa(empresa_sabana, empresas_pista)
                reporte.append({
                    "incluida": incluida,
                    "archivo_sabana": item.filename,
                    "empresa_sabana": empresa_sabana,
                    "empresa_pista": empresa_pista,
                    "motivo": motivo,
                })
                if incluida:
                    zout.writestr(item, zin.read(item.filename))
                    total_filtradas += 1
    reporte_salida = _guardar_reporte_coincidencias(reporte)
    return salida, reporte_salida, total_generadas, total_filtradas


def main() -> int:
    if not ARCHIVO_PISTA.exists():
        raise FileNotFoundError(f"No existe {ARCHIVO_PISTA}")

    app = create_app(os.environ["FLASK_ENV"])
    with app.app_context():
        usuario = _usuario_para_proceso()
        empresas_pista = leer_empresas_pista()
        salida, reporte_salida, total_generadas, total_filtradas = generar_zip_sabanas(app, usuario, empresas_pista)

    print("Proceso terminado.")
    print(f"Periodo: {FECHA_DESDE:%Y-%m-%d} a {FECHA_HASTA:%Y-%m-%d}")
    print(f"Empresas en PISTA.xlsx: {len(empresas_pista)}")
    print(f"Sabanas generadas por el proceso actual antes de filtrar: {total_generadas}")
    print(f"Sabanas coincidentes con PISTA.xlsx: {total_filtradas}")
    print(f"Archivo generado: {salida}")
    print(f"Reporte de coincidencias: {reporte_salida}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
