"""Carga el catalogo retencion_fuente_colombia_2026.xlsx en la base de datos."""
from datetime import date
from openpyxl import load_workbook
from app import create_app
from app.models import RetencionFuente, db

app = create_app()
with app.app_context():
    ws = load_workbook('retencion_fuente_colombia_2026.xlsx', data_only=True)['Retenciones_2026']
    encabezados = [str(x.value).strip() for x in ws[1]]
    creados = actualizados = 0
    for valores in ws.iter_rows(min_row=2, values_only=True):
        fila = dict(zip(encabezados, valores))
        item = RetencionFuente.query.filter_by(codigo=str(fila['codigo']).strip()).first()
        if not item:
            item = RetencionFuente(codigo=str(fila['codigo']).strip()); db.session.add(item); creados += 1
        else: actualizados += 1
        for campo_excel, campo_modelo in [('categoria','categoria'), ('concepto','concepto'), ('base_uvt','base_uvt'), ('base_pesos_2026','base_pesos'), ('tarifa_porcentaje','tarifa_porcentaje'), ('comparador','comparador'), ('aplicacion_automatica','aplicacion_automatica'), ('observaciones','observaciones'), ('norma_referencia','norma_referencia')]:
            setattr(item, campo_modelo, fila.get(campo_excel))
        for campo in ('vigencia_desde', 'vigencia_hasta'):
            valor = fila.get(campo)
            setattr(item, campo, date.fromisoformat(str(valor)) if valor else None)
    db.session.commit()
    print(f'Retenciones cargadas. Nuevas: {creados}; actualizadas: {actualizados}.')
